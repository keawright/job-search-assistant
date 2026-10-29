"""Private SQLite history with all-sheet exclusions and explicit decisions."""
from dataclasses import asdict
from difflib import SequenceMatcher
import json
import re
import sqlite3
from pathlib import Path

from openpyxl import load_workbook
from .normalization import identities, company_key, title_key


class Store:
    def __init__(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(path)
        self.db.executescript('''
            CREATE TABLE IF NOT EXISTS exclusions(identity TEXT PRIMARY KEY, origin TEXT);
            CREATE TABLE IF NOT EXISTS pairs(company TEXT, title TEXT, origin TEXT, PRIMARY KEY(company,title));
            CREATE TABLE IF NOT EXISTS findings(key TEXT PRIMARY KEY, payload TEXT, hash TEXT, state TEXT,
                reason TEXT, decision TEXT, first_seen TEXT, last_seen TEXT, reported_on TEXT);
            CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY, run TEXT, key TEXT, state TEXT, reason TEXT, payload TEXT);
            CREATE TABLE IF NOT EXISTS coverage(id INTEGER PRIMARY KEY, run TEXT, payload TEXT);
            CREATE TABLE IF NOT EXISTS employers(name TEXT PRIMARY KEY, url TEXT, category TEXT, discovered TEXT);
            CREATE TABLE IF NOT EXISTS finding_aliases(identity TEXT,key TEXT,PRIMARY KEY(identity,key));
            CREATE TABLE IF NOT EXISTS finding_pairs(company TEXT,title TEXT,key TEXT,PRIMARY KEY(company,title,key));
        ''')

        # Backfill identities for databases created by earlier search versions.
        for key,payload in self.db.execute('SELECT key,payload FROM findings').fetchall():
            self._index(key,json.loads(payload))
        self.db.commit()

    def _index(self,key,payload):
        for url in [payload['url']] + payload.get('aliases',[]):
            for identity in identities(url):
                self.db.execute('INSERT OR IGNORE INTO finding_aliases VALUES(?,?)',(identity,key))
        company,title=company_key(payload.get('company')),title_key(payload.get('title'))
        if company and title:
            self.db.execute('INSERT OR IGNORE INTO finding_pairs VALUES(?,?,?)',(company,title,key))

    def close(self):
        self.db.commit(); self.db.close()

    def import_workbook(self, path):
        wb = load_workbook(path, read_only=False, data_only=False)
        count = 0
        try:
            for ws in wb:
                headers = None
                for cells in ws.iter_rows():
                    # User-maintained rejection list can be just pasted links,
                    # with no company/title columns or even a header row.
                    if ws.title.strip().casefold() == 'rejected':
                        origin = f'{Path(path).name}:{ws.title}:{cells[0].row}'
                        urls = set()
                        for cell in cells:
                            urls.update(re.findall(r'https?://[^\s<>"\)]+', str(cell.value or '')))
                            if cell.hyperlink and cell.hyperlink.target:
                                urls.add(cell.hyperlink.target)
                        for url in urls:
                            try:
                                for identity in identities(url):
                                    self.db.execute('INSERT OR IGNORE INTO exclusions VALUES(?,?)', (identity, origin))
                            except ValueError:
                                continue
                        if urls: count += 1
                        continue
                    values = [str(c.value or '').strip() for c in cells]
                    if headers is None:
                        if any(v.casefold() in {'company', 'employer'} for v in values) and any('url' in v.casefold() or 'job posting' in v.casefold() for v in values):
                            headers = values
                        continue
                    row = dict(zip(headers, values))
                    company = next((v for k,v in row.items() if k.casefold() in {'company','employer'}), '')
                    title = next((v for k,v in row.items() if k.casefold() in {'position title','job title','title','position'}), '')
                    origin = f'{Path(path).name}:{ws.title}:{cells[0].row}'
                    if company and title and company.casefold() != 'unknown company':
                        self.db.execute('INSERT OR IGNORE INTO pairs VALUES(?,?,?)', (company_key(company), title_key(title), origin))
                    urls = set()
                    for header, cell in zip(headers, cells):
                        if 'url' in header.casefold() or 'job posting' in header.casefold() or 'source link' in header.casefold():
                            urls.update(re.findall(r'https?://[^\s<>"\)]+', str(cell.value or '')))
                        if cell.hyperlink and cell.hyperlink.target:
                            urls.add(cell.hyperlink.target)
                    for url in urls:
                        try:
                            for identity in identities(url):
                                self.db.execute('INSERT OR IGNORE INTO exclusions VALUES(?,?)', (identity, origin))
                        except ValueError:
                            continue
                    if urls or (company and title): count += 1
            self.db.commit()
        finally:
            wb.close()
        return count

    def duplicate(self, posting):
        for url in [posting.url] + posting.aliases:
            for identity in identities(url):
                row = self.db.execute('SELECT origin FROM exclusions WHERE identity=?', (identity,)).fetchone()
                if row: return 'duplicate', 'Already tracked: ' + row[0]
        company, title = company_key(posting.company), title_key(posting.title)
        if company and title:
            for old_title, origin in self.db.execute('SELECT title,origin FROM pairs WHERE company=?', (company,)):
                if title == old_title: return 'duplicate', 'Company/title already tracked: ' + origin
                if SequenceMatcher(None, title, old_title).ratio() >= .86:
                    return 'uncertain', 'Possible company/title duplicate: ' + origin
        for url in [posting.url]+posting.aliases:
            for identity in identities(url):
                if self.db.execute('SELECT 1 FROM finding_aliases WHERE identity=? AND key<>?',(identity,posting.key)).fetchone():
                    return 'duplicate','Previously discovered alternate posting URL'
        if company and title:
            for old_title, in self.db.execute('SELECT title FROM finding_pairs WHERE company=? AND key<>?',(company,posting.key)):
                if title==old_title: return 'duplicate','Previously discovered company/title'
                if SequenceMatcher(None,title,old_title).ratio()>=.86:
                    return 'uncertain','Possible previously discovered company/title duplicate'
        row = self.db.execute('SELECT state,reported_on,reason,decision FROM findings WHERE key=?', (posting.key,)).fetchone()
        if row and (row[1] or row[0] in {'rejected','duplicate','superseded'}):
            detail=json.loads(row[3] or '{}').get('requirements',{}).get('reason','')
            return 'duplicate', 'Previously processed: ' + row[0] + '; ' + row[2] + ('; '+detail if detail else '')
        return '', ''

    def record(self, posting, state, reason, decision, run):
        payload = json.dumps(asdict(posting), ensure_ascii=False)
        previous=self.db.execute('SELECT state,reported_on FROM findings WHERE key=?',(posting.key,)).fetchone()
        if state=='duplicate' and previous and (previous[1] or previous[0] in {'rejected','superseded'}):
            self.db.execute('UPDATE findings SET last_seen=? WHERE key=?',(run,posting.key))
            self.db.execute('INSERT INTO events(run,key,state,reason,payload) VALUES(?,?,?,?,?)',
                            (run,posting.key,state,reason,json.dumps(decision)))
            self.db.commit()
            return

        self.db.execute('''INSERT INTO findings VALUES(?,?,?,?,?,?,?,?,NULL)
            ON CONFLICT(key) DO UPDATE SET payload=excluded.payload,hash=excluded.hash,state=excluded.state,
            reason=excluded.reason,decision=excluded.decision,last_seen=excluded.last_seen''',
            (posting.key, payload, posting.content_hash, state, reason, json.dumps(decision), posting.discovered, run))
        self._index(posting.key,asdict(posting))
        self.db.execute('INSERT INTO events(run,key,state,reason,payload) VALUES(?,?,?,?,?)',
                        (run, posting.key, state, reason, json.dumps(decision)))
        self.db.commit()

    def coverage(self, run, payload):
        self.db.execute('INSERT INTO coverage(run,payload) VALUES(?,?)', (run, json.dumps(payload)))
        self.db.commit()
