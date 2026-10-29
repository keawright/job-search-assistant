"""Recheck saved leads; never edits the application tracker or bypasses fit gates."""
from dataclasses import asdict
from datetime import date,datetime
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from ..archive import current_workbook,fingerprint
from ..assessment import load_service
from .normalization import Posting,canonical,identities,company_key,title_key
from .storage import Store
from .retrieval import PublicClient
from .verification import inspect_connected,verify_live,current_application_verified,choose_verification
from .evaluation import decide,review_hash
from .reporting import export
from .locking import run_lock
from .eligibility import exclusion_reason


def verify_leads(folder,*,report=None,captures=None,offline=False,limit=100,progress=print):
    root=Path(folder).resolve(); private=root/'.job-search-assistant-private/search'
    with run_lock(private/'.run.lock'):
        return _verify(root,private,report,captures,offline,limit,progress)


def _verify(root,private,report,captures,offline,limit,progress):
    report=Path(report) if report else max((private/'reports').glob('job-search-*.xlsx'),key=lambda p:p.stat().st_mtime)
    tracker=current_workbook(root); before=fingerprint(tracker)
    store=Store(private/'discovery.sqlite3');db=store.db
    run=datetime.now().strftime('%Y%m%d-%H%M%S-%f')+'-verification'; today=date.today().isoformat()
    try:
        store.import_workbook(tracker)
        db.execute('''CREATE TABLE IF NOT EXISTS application_checks(
            id INTEGER PRIMARY KEY,run TEXT,key TEXT,checked_on TEXT,status TEXT,url TEXT,method TEXT,evidence TEXT,capture TEXT)''')
        service=load_service(root,offline=True)
        resume=service.resume if service else ''
        additions=private/'candidate-additions.md'
        if additions.exists():resume+='\n\n'+additions.read_text(encoding='utf-8')
        reviews={}
        bundle=private/'reviews.json'
        if bundle.exists():reviews={canonical(i['posting']['url']):i for i in json.loads(bundle.read_text(encoding='utf-8'))}
        captured={}; coverage=[]
        if captures:
            for path in sorted(Path(captures).glob('lead-*.json')):
                c=json.loads(path.read_text(encoding='utf-8'))
                captured.setdefault(canonical(c['url']),[]).append((path,c))
        wb=load_workbook(report,data_only=True);input_rows=[]
        try:
            for name,state in [('Recommended Jobs','recommended'),('Needs Verification','verification'),('Rejected Jobs','rejected')]:
                ws=wb[name];headers=[c.value for c in ws[1]]
                for cells in ws.iter_rows(min_row=2,values_only=True):input_rows.append((dict(zip(headers,cells)),state))
        finally:wb.close()
        client=PublicClient(budget=max(1,limit*2));attempted=0;results=[];checked=set()
        for row,state in input_rows:
            url=row.get('Application URL')
            if not url:continue
            stored=db.execute('SELECT payload,decision FROM findings WHERE key=?',(Posting(url).key,)).fetchone()
            p=Posting(**json.loads(stored[0])) if stored else Posting(url,row.get('Company') or '',row.get('Job title') or '',source=row.get('Evidence source') or '')
            review=json.loads(stored[1]) if stored else {}
            reason=row.get('Decision reason') or ''
            hard_exclusion=exclusion_reason(p)
            if state in {'recommended','verification'} and hard_exclusion:
                store.record(p,'rejected',hard_exclusion,review,run)
                results.append((p,'rejected',hard_exclusion,review))
                continue
            if state in {'recommended','verification'}:
                proof_candidates=list(captured.get(canonical(p.url),[]))
                # An explicitly linked /apply page is the same vacancy by ATS ID.
                for capture_url,items in captured.items():
                    if capture_url!=canonical(p.url) and identities(capture_url)&identities(p.url):proof_candidates+=items
                checks=[]
                for path,c in proof_candidates:
                    probe=Posting(**asdict(p))
                    inspect_connected(probe,c['text'],c.get('checked_on',today))
                    checks.append(probe);checked.add(p.key)
                    db.execute('INSERT INTO application_checks(run,key,checked_on,status,url,method,evidence,capture) VALUES(?,?,?,?,?,?,?,?)',
                               (run,p.key,probe.application_checked_on,probe.application_status,probe.application_verification_url,
                                probe.application_verification_method,probe.application_verification_evidence,str(path)))
                    coverage.append(dict(source=c['url'],method='Connected application-status retrieval',checked=c.get('checked_on',today),discovered=0,assessed=0,error=probe.application_status+': '+probe.application_verification_evidence))
                if checks:
                    p=choose_verification(checks)
                elif not offline and attempted<limit:
                    attempted+=1;checked.add(p.key)
                    p.application_status='unverified';p.application_verification_fresh=False;p.application_checked_on=today
                    try:verify_live(p,client)
                    except (ValueError,OSError) as exc:p.application_verification_evidence='Current status check failed: '+str(exc)
                    db.execute('INSERT INTO application_checks(run,key,checked_on,status,url,method,evidence,capture) VALUES(?,?,?,?,?,?,?,?)',
                               (run,p.key,p.application_checked_on,p.application_status,p.application_verification_url,
                                p.application_verification_method,p.application_verification_evidence,''))
                    coverage.append(dict(source=p.url,method='Direct application-status check',checked=today,discovered=0,assessed=0,error=p.application_status+': '+p.application_verification_evidence))
                # A valid fit assessment may be reused only for unchanged inputs.
                item=reviews.get(canonical(p.url))
                if item and item.get('input_hash')==review_hash(p,resume): review=item['review']
                if p.application_status=='closed':state,reason='rejected','Employer reports applications closed: '+p.application_verification_evidence
                elif not current_application_verified(p):
                    state='verification'
                    reason=('Verification is expired or not bound to this vacancy. ' if p.application_status=='accepting' else '')+(p.application_verification_evidence or 'Current application acceptance has not been verified today; no current retrieval within this run budget')
                elif review:
                    try:state,reason=decide(p,review,resume)
                    except (ValueError,TypeError,KeyError):state,reason='verification','Application acceptance verified; current grounded fit review still required'
                else:state,reason='verification','Application acceptance verified; full fit/location/mandatory-requirement review still required'
                # Recheck current tracker exclusions independently of discovery's own history.
                old=[r[0] for u in [p.url]+p.aliases for identity in identities(u) for r in db.execute('SELECT origin FROM exclusions WHERE identity=?',(identity,))]
                pair=db.execute('SELECT origin FROM pairs WHERE company=? AND title=?',(company_key(p.company),title_key(p.title))).fetchone()
                if old or pair:state,reason='duplicate','Already tracked: '+(old[0] if old else pair[0])
                if re.search(r'/careers/?(?:[?#].*)?$|/jobs/?(?:[?#].*)?$|/company/[^/]+/jobs|/role/l/|indeed\.com/[a-z]-.*jobs\.html|/roles/|/Jobs/[^/]+/-in-',p.url,re.I):
                    state,reason='rejected','Search/career directory, not a specific application vacancy'
                if state=='verification' and not current_application_verified(p):
                    # Withdraw an earlier unverified recommendation while retaining events.
                    db.execute('UPDATE findings SET reported_on=NULL WHERE key=? AND state=?',(p.key,'recommended'))
                store.record(p,state,reason,review,run)
            results.append((p,state,reason,review))
        if fingerprint(tracker)!=before:raise ValueError('Tracker changed during verification; re-import latest workbook before reporting')
        coverage.append(dict(source='Verification policy',method='Independent same-day employer application form required',checked=today,discovered=len(checked),assessed=sum(current_application_verified(r[0]) for r in results),error='No AI verdict, stale cache, third-party Apply link or generic careers page can establish current acceptance.'))
        outputs=export(private/'reports',run,results,coverage,title='Application status verification')
        summary=dict(run=run,input_report=str(report),tracker_sha256=before,leads_with_retrieval=len(checked),
                     page_retrievals=len([c for c in coverage if c['method']=='Connected application-status retrieval']),
                     accepting_verified=sum(current_application_verified(r[0]) for r in results),
                     recommendations=sum(r[1]=='recommended' for r in results),
                     unresolved=sum(r[1]=='verification' for r in results),outputs=[str(p) for p in outputs])
        (private/f'{run}.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
        for c in coverage:store.coverage(run,c)
        progress(json.dumps(summary,indent=2));return outputs
    finally:store.close()
