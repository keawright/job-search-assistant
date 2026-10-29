"""Private copy-on-write search reports, never edits the application tracker."""
from pathlib import Path
import csv
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment

HEADERS=['Company','Job title','Location','Remote/hybrid/on-site','Salary range','Job posting date',
         'Date discovered','Fit classification','Fit score','Why I am a good fit','Potential concerns',
         'Application URL','Application status','Decision reason','Work-life balance','Evidence source',
         'Posting application status','Application status checked on','Verification URL','Verification method','Verification evidence']


def safe(value):
    if isinstance(value,(int,float)) and not isinstance(value,bool): return value
    value=ILLEGAL_CHARACTERS_RE.sub('',str(value or ''))[:32767]
    return "'"+value if value.startswith(('=','+','-','@')) else value


def row(posting,state,reason,review):
    a=review.get('assessment',{})
    classification=reason if state=='recommended' else ('Stretch' if any(a.get(k,{}).get('rating')=='stretch' for k in ['leadership_fit','technical_fit','career_alignment','opportunity_quality']) else 'Not recommended')
    return [posting.company,posting.title,posting.location,posting.arrangement,posting.salary or review.get('compensation',''),
            posting.posted,posting.discovered,classification,4 if classification=='Strong Fit' else 3 if classification=='Good Fit' else '',
            a.get('summary',''), '\n'.join(a.get('uncertainties',[])+[g.get('description','') for g in a.get('gaps',[])]),
            posting.url,'Not applied',reason,review.get('work_life_balance','Unknown'),posting.source,
            posting.application_status,posting.application_checked_on,posting.application_verification_url,
            posting.application_verification_method,posting.application_verification_evidence]


def export(folder,run,results,coverage,*,title="New job opportunities"):
    from .eligibility import exclusion_reason
    results=[(p,'rejected',exclusion_reason(p),review)
             if state in {'recommended','verification'} and exclusion_reason(p) else (p,state,reason,review)
             for p,state,reason,review in results]
    from .verification import current_application_verified, current_closure_verified
    results=[(p,'rejected','Employer vacancy closed or removed: '+p.application_verification_evidence,review)
             if state in {'recommended','verification'} and current_closure_verified(p) else (p,state,reason,review)
             for p,state,reason,review in results]
    results=[(p,'verification','Recommendation withheld: current application acceptance is not verified today',review)
             if state=='recommended' and not current_application_verified(p) else (p,state,reason,review)
             for p,state,reason,review in results]
    folder=Path(folder); folder.mkdir(parents=True,exist_ok=True)
    xlsx=folder/f'job-search-{run}.xlsx'; md=folder/f'job-search-{run}.md'
    if xlsx.exists() or md.exists(): raise FileExistsError('Search report already exists')
    wb=Workbook(); wb.remove(wb.active)
    groups={'Recommended Jobs':{'recommended'},'Needs Verification':{'verification'},
            'Rejected Jobs':{'rejected','duplicate','superseded'}}
    for name,states in groups.items():
        ws=wb.create_sheet(name); ws.append(HEADERS)
        for p,state,reason,review in results:
            if state in states:
                ws.append([safe(x) for x in row(p,state,reason,review)])
                if p.url.startswith(('https://','http://')):
                    ws.cell(ws.max_row,12).hyperlink=p.url
                    ws.cell(ws.max_row,12).style='Hyperlink'
    ws=wb.create_sheet('Search Coverage')
    ws.append(['Source','Search method','Date checked','Number of postings discovered','Number assessed','Errors or access limitations'])
    for c in coverage: ws.append([safe(c.get(k,'')) for k in ['source','method','checked','discovered','assessed','error']])
    for ws in wb:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='24445C')
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=32
        for cells in ws.iter_rows(min_row=2):
            for cell in cells: cell.alignment=Alignment(vertical='top',wrap_text=True)
        if ws.title!='Search Coverage':
            for column in ['J','K','N','O']: ws.column_dimensions[column].width=65
            ws.column_dimensions['B'].width=48
            ws.column_dimensions['L'].width=55
    with xlsx.open('xb') as stream: wb.save(stream)
    wb.close()
    for name in groups:
        with (folder/f'{xlsx.stem}-{name.lower().replace(" ","-")}.csv').open('x',encoding='utf-8-sig',newline='') as stream:
            writer=csv.writer(stream); writer.writerow(HEADERS)
            for p,state,reason,review in results:
                if state in groups[name]: writer.writerow([safe(x) for x in row(p,state,reason,review)])
    recommendations=[r for r in results if r[1]=='recommended']
    lines=['# '+title,f'Run: {run}',
           f'{len(recommendations)} recommendations; {sum(r[1]=="verification" for r in results)} need verification.',
           'Fit score is an ordinal label (Strong=4, Good=3), not a hiring probability.',
           'Posting dates remain blank when unknown; crawl/update dates are not original posting dates.','']
    for i,(p,state,reason,review) in enumerate(recommendations[:20],1):
        lines += [f'## {i}. {p.company} — {p.title}',f'[{reason}]({p.url})',review['assessment']['summary'],
                  'Concerns: '+ '; '.join(review['assessment']['uncertainties']), '']
    if not recommendations: lines += ['No new positions passed every recommendation gate. Verification candidates are not recommendations.','']
    verification=[r for r in results if r[1]=='verification' and r[3].get('assessment')]
    if verification:
        lines += ['## Worth verifying, not recommendations',
                  'These roles have useful fit evidence but fail one or more required recommendation gates.']
        for p,state,reason,review in verification[:10]:
            lines += [f'- [{p.company} - {p.title}]({p.url}): {reason}. '+review['assessment']['summary']]
        lines.append('')
    lines += ['## Coverage and limitations']
    connected=[c for c in coverage if c['method'].startswith('Connected search:')]
    queries=sum(c['method'].count(' | ')+1 for c in connected)
    pages=[c for c in coverage if c['method']=='Connected public page retrieval']
    skipped=sum(str(c.get('error','')).startswith('Not executed:') for c in coverage)
    lines += [f'- {queries} connected-search queries across {len(connected)} saved batches.',
              f'- {len(pages)} public page/application retrievals recorded.',
              f'- {skipped} planned direct queries skipped after the local search host failed.',
              '- Detailed per-source counts and errors are in the Search Coverage worksheet.',
              '- Search-result URLs include non-vacancy pages and documents, not just jobs.',
              '- Missing/ambiguous location, hard requirements and current application status remain holds.']
    errors=list(dict.fromkeys(str(c.get('error','')) for c in coverage if c.get('error') and not str(c.get('error','')).startswith(('Not executed:','Search-index evidence','Public page read'))))
    for error in errors[:8]: lines.append('- '+error[:300])
    md.write_text('\n'.join(lines),encoding='utf-8')
    return xlsx,md
