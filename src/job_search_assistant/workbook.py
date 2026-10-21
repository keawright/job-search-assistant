"""Copy-on-write workbook workflow with explicit ownership of generated fields."""

from copy import copy
from calendar import monthrange
from dataclasses import asdict, dataclass, field, replace
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from .ingestion import FetchError, closure_evidence, extract_posting, fetch_page, normalize_url
from .discovery import Identity, find_alternative, is_linkedin, linkedin_identity, search_postings
from .models import Assessment, CandidatePreferences, Gap, Job, Location
from .ranking import prioritize
from .assessment import CRITERIA, AssessmentError

QUEUE_SHEET = "Assistant queue"
MANUAL_COLUMNS = ["Manual priority", "Leadership fit", "Technical fit", "Career alignment", "Opportunity quality",
                  "Assessment evidence", "Assessment strengths", "Assessment gaps (JSON)", "Regular hands-on coding", "Coding evidence",
                  "Working arrangement", "Remote countries", "Office city", "Relocation required", "Required work authorization",
                  "Location evidence", "Verification status", "Last verified"]
GENERATED_COLUMNS = ["Apply order", "Assistant ID", "Assistant section", "Assistant priority", "Assistant overall fit", "Assistant location eligibility",
                     "Assistant priority reasons", "Assistant review flags", "Assistant evaluated on", "Assistant policy version",
                     "Assistant fetched on", "Assistant resolved URL", "Assistant extracted posting date", "Assistant extracted deadline",
                     "Assistant extracted location", "Assistant description", "Assistant snapshot", "Assistant extraction flags", "Assistant fetch issue",
                     "Assistant availability checked on", "Assistant closure evidence",
                     "Assistant alternate URL", "Assistant discovery notes", "Assistant discovery identity", "Assistant assessment status",
                     "Assistant assessment source", "Assistant assessed on", "Assistant assessment summary", "Assistant assessment evidence",
                     "Assistant assessment strengths", "Assistant assessment gaps", "Assistant assessment uncertainties", "Assistant assessment issue"]
ALIASES = {"company": ("Company",), "title": ("Position title", "Position", "Title"),
           "url": ("Job posting (URL)", "Job posting", "Job posting URL"), "status": ("Application status",),
           "deadline": ("Deadline", "Application deadline"), "posting_date": ("Posted date", "Posting date")}


@dataclass
class Record:
    values: dict
    row: int | None = None
    flags: list[str] = field(default_factory=list)
    job: Job | None = None


@dataclass(frozen=True)
class UpdateSummary:
    output: Path
    source_sheet: str
    imported: int
    added: int
    skipped_urls: int
    fetch_failures: int
    outstanding: int
    review: int


def _column(headers, key):
    return next((h for h in ALIASES[key] if h in headers), None)


def _value(record, key):
    return next((record.values[h] for h in ALIASES[key] if record.values.get(h) is not None), None)


def _day(value, epoch):
    if value is None or str(value).strip().casefold() in ("", "unknown", "not stated", "—"):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return from_excel(value, epoch).date()
    try:
        return date.fromisoformat(str(value).strip()[:10])
    except ValueError:
        # Do not guess ambiguous numeric locale dates.
        return None


def _bool(value):
    if value is None or str(value).strip().lower() in ("", "unknown", "—"):
        return None
    if str(value).strip().lower() in ("true", "yes"):
        return True
    if str(value).strip().lower() in ("false", "no"):
        return False
    raise ValueError(f"Expected yes/no/unknown, got {value!r}")


def _urls(value):
    return re.findall(r"https?://[^\s<>]+", str(value or ""))


def _source_url(cell):
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    if cell.data_type == "f" and str(cell.value).upper().startswith("=HYPERLINK("):
        match = re.match(r'=HYPERLINK\("([^"]+)"', str(cell.value), re.I)
        return match.group(1) if match else None
    return cell.value


def _read_records(wb, requested_sheet):
    sheet_name = requested_sheet or next((name for name in (QUEUE_SHEET, "Updated priorities", "Application priorities") if name in wb.sheetnames), None)
    if sheet_name not in wb.sheetnames:
        raise ValueError("Choose an existing job table with --sheet")
    ws = wb[sheet_name]
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        labels = [str(c.value).strip() if c.value is not None else "" for c in row]
        if _column(labels, "company") and _column(labels, "title") and _column(labels, "status") and _column(labels, "url"):
            header_row = row[0].row
            break
    if header_row is None:
        raise ValueError("Cannot find Company, Position title, Job posting, and Application status headers")
    headers = [str(c.value).strip() if c.value is not None else f"Source column {c.column}" for c in ws[header_row]]
    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate column headers require explicit reconciliation")
    records = []
    for row in ws.iter_rows(min_row=header_row + 1):
        values = {h: c.value for h, c in zip(headers, row)}
        record = Record(values, row[0].row)
        if not _value(record, "company") and not _value(record, "title"):
            continue
        for h, c in zip(headers, row):
            if h in ALIASES["url"]:
                # The original cell itself is copied unchanged during export.
                record.values["_posting_url"] = _source_url(c)
            elif h in ("Other source links", "Other source links (URL)") and c.hyperlink:
                record.values["_other_url"] = c.hyperlink.target
        record.values.setdefault("Assistant ID", None)
        if not record.values["Assistant ID"]:
            record.values["Assistant ID"] = f"import-{len(records)+1:06d}"
        records.append(record)
    return ws, headers, records


def _assessment(values):
    criteria = {key: str(values.get(label) or "unknown").strip().lower()
                for key, label in (("leadership_fit", "Leadership fit"), ("technical_fit", "Technical fit"),
                                   ("career_alignment", "Career alignment"), ("opportunity_quality", "Opportunity quality"))}
    evidence = values.get("Assessment evidence") or ""
    raw_gaps = values.get("Assessment gaps (JSON)")
    gaps = tuple(Gap(**g) for g in json.loads(raw_gaps)) if raw_gaps else ()
    return Assessment(**criteria, evidence={k: str(evidence) for k in criteria},
                      strengths=tuple(str(values.get("Assessment strengths") or "").splitlines()), gaps=gaps,
                      regular_hands_on_coding=_bool(values.get("Regular hands-on coding")), coding_evidence=values.get("Coding evidence"))


def _location(values):
    human_fields = ("Working arrangement", "Remote countries", "Office city", "Relocation required", "Required work authorization", "Location evidence")
    if any(values.get(h) is not None and str(values[h]).strip() for h in human_fields):
        # Human values take complete precedence, avoiding mixtures of contradictory sources.
        countries = values.get("Remote countries")
        return Location(arrangement=str(values.get("Working arrangement") or "unknown").strip().lower(),
                        allowed_remote_countries=tuple(c.strip() for c in str(countries).split(",") if c.strip()) if countries else None,
                        office_city=values.get("Office city"), relocation_required=_bool(values.get("Relocation required")),
                        required_work_authorization=values.get("Required work authorization"), evidence=values.get("Location evidence"))
    raw = values.get("Assistant extracted location")
    if raw:
        data = json.loads(raw)
        if data.get("allowed_remote_countries") is not None:
            data["allowed_remote_countries"] = tuple(data["allowed_remote_countries"])
        return Location(**data)
    return Location()


def _make_job(record, epoch):
    v = record.values
    raw_status = str(_value(record, "status") or "").strip().casefold()
    if raw_status == "closed":
        raise ValueError("Closed job retained outside the application queue; not checked again automatically")
    status = {"applied": "applied", "not applied": "outstanding", "outstanding": "outstanding", "": "outstanding", "rejected": "rejected", "elapsed-rejected": "rejected"}.get(raw_status)
    if status is None:
        raise ValueError(f"Unrecognized application status {raw_status!r}; preserved for review")
    if not raw_status:
        record.flags.append("Blank application status treated as Outstanding; original cell preserved.")
    url = v.get("_posting_url") or _value(record, "url")
    if not url:
        raise ValueError("Missing posting URL; opportunity retained for review")
    normalize_url(str(url))
    verification = str(v.get("Verification status") or "unverified").strip().casefold()
    legacy_vacancy = str(v.get("Vacancy status") or "").strip().casefold()
    future_cgi = str(_value(record, "company") or "").strip().casefold() == "cgi" and legacy_vacancy.startswith("future opportunity")
    if future_cgi:
        url = v.get("Assistant alternate URL") or url
        normalize_url(str(url))
        record.flags.append("Future opportunity: fit assessed at user request; a current staffed vacancy is not confirmed.")
    if not future_cgi and not v.get("Verification status") and (legacy_vacancy.startswith(("closed", "apparently closed", "future opportunity", "open application — no current")) or legacy_vacancy == "deadline passed"):
        # Preserve the existing exclusion, but do not invent a verification date.
        raise ValueError("Legacy vacancy status indicates closed/passed/no current vacancy; verify before returning to the queue")
    legacy_location = str(v.get("Can work from Halifax?") or "").strip().casefold()
    if legacy_location.startswith(("no ", "remote us only", "remote uk only")):
        record.flags.append("Location needs checking; assessed and ranked assuming location will be acceptable. Prior location note: " + str(v.get("Can work from Halifax?")))
    manual = v.get("Manual priority")
    if manual is not None and str(manual).strip():
        number = float(manual)
        if not number.is_integer():
            raise ValueError("Manual priority must be a positive integer")
        manual = int(number)
    else:
        manual = None
    for key in ("deadline", "posting_date"):
        raw = _value(record, key)
        if raw is not None and _day(raw, epoch) is None:
            record.flags.append(f"Unrecognized {key} value retained; date treated as unknown.")
    return Job(id=str(v["Assistant ID"]), company=str(_value(record, "company") or "Unknown company"),
               title=str(_value(record, "title") or "Unknown position"), url=str(url), status=status,
               verification=verification, last_verified=_day(v.get("Last verified"), epoch),
               posting_date=_day(_value(record, "posting_date"), epoch) or _day(v.get("Assistant extracted posting date"), epoch),
               deadline=_day(_value(record, "deadline"), epoch) or _day(v.get("Assistant extracted deadline"), epoch),
               description=v.get("Assistant description"), location=_location(v), assessment=_assessment(v), manual_priority=manual,
               assume_location_acceptable=True,
               source_urls=tuple(_urls(v.get("_other_url") or v.get("Other source links (URL)") or v.get("Other source links"))),
               tracking={k: value for k, value in v.items() if k not in GENERATED_COLUMNS and not k.startswith("_")})


def _age_applied(record, epoch, as_of):
    if str(_value(record, "status") or "").strip().casefold() != "applied":
        return
    raw = next((value for label, value in record.values.items()
                if label.strip().casefold() in ("applied date", "date applied", "application date")
                and value is not None), None)
    try:
        applied = _day(raw, epoch)
    except (ValueError, TypeError, OverflowError, AttributeError):
        applied = None
    if applied is None and isinstance(raw, str):
        for fmt in ("%b-%d-%Y", "%B-%d-%Y", "%b %d, %Y", "%B %d, %Y"):
            try:
                applied = datetime.strptime(raw.strip(), fmt).date()
                break
            except ValueError:
                pass
    if applied is None or applied > as_of:
        record.flags.append("Applied aging skipped: application date is missing, unreadable, or in the future; status retained.")
        return
    if applied.year == 9999 and applied.month == 12:
        return
    year = applied.year + (applied.month == 12)
    month = applied.month % 12 + 1
    anniversary = date(year, month, min(applied.day, monthrange(year, month)[1]))
    if as_of >= anniversary:
        record.values[_column(record.values, "status")] = "Elapsed-Rejected"


def _usable_posting(posting):
    return bool(posting and posting.company and posting.title and posting.description
                and not any("No JobPosting metadata" in flag or "Multiple JobPosting" in flag for flag in posting.flags))


def _supplied_alternative(record, extracted, current_url, fetcher):
    """Use only a structured, substantive posting matching the known job identity."""
    def normalized(value):
        return " ".join(re.findall(r"\w+", str(value or "").casefold()))

    company = (extracted.company if extracted else None) or _value(record, "company")
    title = (extracted.title if extracted and extracted.company else None) or _value(record, "title")
    identity = linkedin_identity(str(title or ""))
    if identity:
        company, title = identity.company, identity.title
    if not company or not title or any(word in normalized(company) + " " + normalized(title)
                                     for word in ("unknown", "unidentified", "not verified", "needs posting review", "undisclosed")):
        return None, "Supplied alternate links need review: employer/role identity is insufficient to verify a match."
    urls = [record.values.get("_other_url")]
    urls += _urls(record.values.get("Other source links (URL)"))
    urls += _urls(record.values.get("Other source links"))
    seen = {normalize_url(current_url)}
    notes = []
    for url in urls:
        if not url:
            continue
        try:
            key = normalize_url(url)
            if key in seen:
                continue
            seen.add(key)
            if len(seen) > 6:
                notes.append("Additional supplied links require review (five-link limit).")
                break
            page = fetcher(url)
            posting = extract_posting(page.html)
            if (not _usable_posting(posting) or normalized(posting.company) != normalized(company)
                    or normalized(posting.title) != normalized(title)):
                notes.append(f"Supplied alternate did not verify the same employer/role with a description: {url}")
                continue
            return page, f"Supplied alternate matched structured employer and role: {page.url}. Original posting URL retained; location and source authority still require review."
        except (FetchError, ValueError):
            notes.append(f"Supplied alternate could not be read: {url}")
    return None, "\n".join(notes)


def _fetch(record, url, snapshots, as_of, fetcher, searcher=search_postings):
    eligible = str(_value(record, "status") or "").strip().casefold() in ("", "not applied")
    original_url = url
    url = record.values.get("Assistant alternate URL") or url
    if eligible:
        record.values["Assistant availability checked on"] = as_of
    try:
        page = fetcher(url)
        extracted = extract_posting(page.html)
    except (FetchError, ValueError) as exc:
        record.values["Assistant fetch issue"] = f"URL read failed: {exc}. Existing information retained; availability unverified."
        page = extracted = None
    original_page = page
    if eligible and not _usable_posting(extracted) and not (page and closure_evidence(page.html)):
        alternate, notes = _supplied_alternative(record, extracted, url, fetcher)
        if notes:
            record.values["Assistant discovery notes"] = notes
        if alternate:
            page = alternate
            extracted = extract_posting(page.html)
            record.values["Assistant alternate URL"] = page.url
    if eligible and is_linkedin(original_url) and not record.values.get("Assistant alternate URL"):
        identity = linkedin_identity((extracted.title if extracted else None) or str(_value(record, "title") or ""))
        if identity is None and record.values.get("Assistant discovery identity"):
            try:
                identity = Identity(**json.loads(record.values["Assistant discovery identity"]))
            except (ValueError, TypeError):
                pass
        if identity and (extracted is None or not extracted.company):
            record.values["Assistant discovery identity"] = json.dumps(asdict(identity), ensure_ascii=False)
            # Repair only generated placeholders and the previous fallback page
            # title. Never replace a manually entered company or clean title.
            company_column, title_column = _column(record.values, "company"), _column(record.values, "title")
            if _value(record, "company") in (None, "Unknown company"):
                record.values[company_column] = identity.company
            current_title = str(_value(record, "title") or "")
            if record.row is None or current_title == "Needs posting review" or linkedin_identity(current_title):
                record.values[title_column] = identity.title
            try:
                alternate = find_alternative(identity, fetcher, searcher)
                if alternate:
                    page = alternate
                    extracted = extract_posting(page.html)
                    record.values["Assistant alternate URL"] = page.url
                    record.values["Assistant discovery notes"] = "Alternate posting matched structured company/title and visible location. Source authority still needs review. Original posting URL retained."
                else:
                    record.values["Assistant discovery notes"] = "No single matching alternate posting verified; search may be blocked, empty, or ambiguous. Original source retained for review."
            except (FetchError, ValueError) as exc:
                record.values["Assistant discovery notes"] = f"Alternate search could not be completed: {exc}. Original source retained for review."
        elif extracted is None or not extracted.company:
            record.values["Assistant discovery notes"] = "Cannot search safely: LinkedIn company and role could not be extracted from the available title."
    if page is None:
        return False
    snapshots.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256((url + page.html).encode()).hexdigest()
    snapshot = snapshots / f"{digest}.json"
    payload = {"requested_url": url, "resolved_url": page.url, "fetched_on": as_of.isoformat(), "html": page.html, "extracted": asdict(extracted),
               "original_html": original_page.html if original_page else None, "original_url": original_url}
    if not snapshot.exists():
        with snapshot.open("x", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, default=str, indent=2)
    v = record.values
    v.update({"Assistant fetched on": as_of, "Assistant resolved URL": page.url,
              "Assistant extracted posting date": extracted.posting_date, "Assistant extracted deadline": extracted.deadline,
              "Assistant extracted location": json.dumps(asdict(extracted.location), ensure_ascii=False),
              "Assistant description": (extracted.description or "")[:30000], "Assistant snapshot": str(snapshot.resolve()),
              "Assistant extraction flags": "\n".join(extracted.flags), "Assistant fetch issue": None})
    if eligible:
        evidence = closure_evidence(page.html)
        if evidence:
            v[_column(v, "status")] = "Closed"
            v["Assistant closure evidence"] = f"{page.url}: {evidence}"
        else:
            v["Assistant fetch issue"] = "No explicit closure notice found; availability remains unverified."
    if extracted.description and len(extracted.description) > 30000:
        record.flags.append("Description shortened in Excel; complete source and text are in the private snapshot.")
    # These fields belong to the user once a record exists. Only fill new rows.
    if record.row is None:
        if extracted.company:
            v[_column(v, "company")] = extracted.company
        if extracted.title:
            identity = linkedin_identity(extracted.title) if is_linkedin(original_url) else None
            v[_column(v, "title")] = identity.title if identity else extracted.title
    return True


def _write_text(cell, value):
    """Untrusted page text must never become an Excel formula."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def update_workbook(source: Path, output: Path, urls: list[str], *, as_of: date,
                    sheet: str | None = None, candidate: CandidatePreferences | None = None,
                    offline: bool = False, refresh: bool = False, compact: bool = False, fetcher=fetch_page,
                    assessment_service=None, reassess_rows: set[int] | None = None, reassess_needed: bool = False) -> UpdateSummary:
    source, output = Path(source).resolve(), Path(output).resolve()
    if source.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
        raise ValueError("Input and output must be .xlsx workbooks (macros are not supported)")
    if source == output or output.exists():
        raise ValueError("Choose a new output filename; existing files are never overwritten")
    # Validate the entire URL batch before fetching or creating files.
    normalized_inputs = [(url.strip(), normalize_url(url)) for url in urls if url.strip() and not url.lstrip().startswith("#")]
    candidate = candidate or CandidatePreferences()
    wb = load_workbook(source, rich_text=True)
    try:
        ws, headers, records = _read_records(wb, sheet)
        if reassess_needed:
            if reassess_rows is not None:
                raise ValueError("Choose reassessment by rows or by Apply order status, not both")
            reassess_rows = {record.row for record in records
                if " ".join(str(record.values.get("Apply order") or "").split()).casefold()
                   in ("needs assessment", "needs assessement", "needs review")
                and str(_value(record, "status") or "").strip().casefold() in ("", "not applied")}
        if reassess_rows is not None:
            if offline or assessment_service is None or not assessment_service.allow_network or assessment_service.provider is None:
                raise ValueError("Selected reassessment requires an enabled AI provider; omit --offline/--no-ai")
            selected = {r.row: r for r in records}
            for row in reassess_rows:
                if row not in selected:
                    raise ValueError(f"Reassessment row {row} is not a job row in {ws.title}")
                if str(_value(selected[row], "status") or "").strip().casefold() not in ("", "not applied"):
                    raise ValueError(f"Reassessment row {row} is protected by its application status")
        source_sheet_name = ws.title
        imported = len(records)
        for record in records:
            _age_applied(record, wb.epoch, as_of)
        seen = {}
        for record in records:
            candidates = [record.values.get("_posting_url"), record.values.get("Assistant resolved URL"), record.values.get("_other_url")]
            candidates += _urls(record.values.get("Other source links (URL)") or record.values.get("Other source links"))
            for url in candidates:
                if url:
                    try:
                        seen.setdefault(normalize_url(str(url)), record)
                    except ValueError:
                        pass
        failures = skipped = added = 0
        snapshots = output.parent / ".job-search-assistant-private" / "snapshots"
        if not offline:
            for record in records:
                if reassess_rows is not None and record.row not in reassess_rows:
                    continue
                if str(_value(record, "status") or "").strip().casefold() not in ("", "not applied"):
                    continue
                url = record.values.get("_posting_url")
                if url:
                    failures += not _fetch(record, str(url), snapshots, as_of, fetcher)
                    resolved = record.values.get("Assistant resolved URL")
                    if resolved:
                        seen.setdefault(normalize_url(resolved), record)
                else:
                    record.values["Assistant availability checked on"] = as_of
                    record.values["Assistant fetch issue"] = "Could not check availability: missing posting URL."
        for url, normalized in normalized_inputs:
            if normalized in seen:
                skipped += 1
                seen[normalized].flags.append("Repeated supplied URL skipped; existing tracking retained.")
                continue
            record = Record({h: None for h in headers})
            record.values.update({_column(headers, "company"): "Unknown company", _column(headers, "title"): "Needs posting review",
                                  _column(headers, "url"): url, _column(headers, "status"): "Not applied",
                                  "Assistant ID": "url-" + hashlib.sha256(normalized.encode()).hexdigest()[:20], "_posting_url": url})
            if offline:
                record.flags.append("Offline: URL has not been fetched; contextual assessment required.")
            else:
                failures += not _fetch(record, url, snapshots, as_of, fetcher)
            resolved = record.values.get("Assistant resolved URL")
            if resolved and normalize_url(resolved) in seen:
                skipped += 1
                existing = seen[normalize_url(resolved)]
                existing.flags.append("A supplied URL resolves to this existing posting; duplicate entry skipped.")
                seen[normalized] = existing
                continue
            records.append(record)
            added += 1
            seen[normalized] = record
            if resolved:
                seen[normalize_url(resolved)] = record
        # Validate one record at a time; malformed user data stays visible.
        valid, ids, job_urls = [], set(), set()
        for record in records:
            if str(_value(record, "status") or "").strip().casefold() == "interviewing":
                continue
            try:
                job = _make_job(record, wb.epoch)
                normalized = normalize_url(job.url)
                if job.id in ids or normalized in job_urls:
                    raise ValueError("Duplicate existing ID/URL; retained separately for manual reconciliation")
                raw_status = str(_value(record, "status") or "").strip().casefold()
                if assessment_service is not None and raw_status in ("", "not applied"):
                    for label in ("Assistant assessment source", "Assistant assessed on", "Assistant assessment summary", "Assistant assessment evidence",
                                  "Assistant assessment strengths", "Assistant assessment gaps", "Assistant assessment uncertainties", "Assistant assessment issue"):
                        record.values[label] = None
                    if any(getattr(job.assessment, key) != "unknown" for key in CRITERIA):
                        record.values["Assistant assessment source"] = "Manual assessment inputs (preserved)"
                    else:
                        try:
                            if reassess_rows is None:
                                assessment, entry = assessment_service.resolve(job, candidate, as_of, offline=offline)
                            else:
                                selected = record.row in reassess_rows
                                assessment, entry = assessment_service.resolve(job, candidate, as_of,
                                    offline=not selected, force=selected)
                            job = replace(job, assessment=assessment)
                            record.values.update({"Assistant assessment source": entry["source"], "Assistant assessed on": entry["assessed_on"],
                                "Assistant assessment summary": entry["assessment"]["summary"],
                                "Assistant assessment evidence": "\n\n".join(f"{key}: {value}" for key, value in assessment.evidence.items()),
                                "Assistant assessment strengths": "\n".join(assessment.strengths),
                                "Assistant assessment gaps": "\n".join(f"{gap.severity}: {gap.description}. {gap.evidence}" for gap in assessment.gaps),
                                "Assistant assessment uncertainties": "\n".join(entry["assessment"]["uncertainties"])})
                        except AssessmentError as exc:
                            record.values["Assistant assessment issue"] = str(exc)
                        except (OSError, ValueError, TypeError, KeyError):
                            # No raw provider/transport details or credentials in workbook output.
                            record.values["Assistant assessment issue"] = "Assessment unavailable: missing key/cache, call limit, insufficient evidence, or invalid response. Run with --no-ai to use saved results; otherwise check AI configuration."
                # Validate date-dependent invariants before batch ranking.
                prioritize([job], candidate, as_of=as_of)
                ids.add(job.id)
                job_urls.add(normalized)
                record.job = job
                valid.append(job)
                if record.row is not None and all(getattr(job.assessment, key) == "unknown" for key in ("leadership_fit", "technical_fit", "career_alignment", "opportunity_quality")):
                    record.flags.append("Legacy fit commentary is preserved but not converted into criterion ratings; supply an evidence-backed assessment for a fresh fit ranking.")
            except (ValueError, TypeError, KeyError, AttributeError, OverflowError) as exc:
                record.flags.append(f"Needs review: {exc}")
        result = prioritize(valid, candidate, as_of=as_of)
        by_id = {r.job.id: r for r in result.ordered + result.excluded}
        record_by_id = {r.job.id: r for r in records if r.job is not None}
        excluded_ids = {r.job.id for r in result.excluded}
        output_records = [record_by_id[r.job.id] for r in result.ordered + result.excluded]
        output_records += [r for r in records if r.job is None]
        for record in output_records:
            v = record.values
            ranked = by_id.get(record.job.id) if record.job else None
            raw_status = str(_value(record, "status") or "").strip().casefold()
            section = {"applied": "Applied", "rejected": "Rejected"}.get(raw_status, "Review")
            if ranked:
                section = "Review" if ranked.job.id in excluded_ids else {"applied": "Applied", "outstanding": "Outstanding", "rejected": "Rejected"}[ranked.job.status]
            if raw_status == "interviewing":
                section = "Interviewing"
            if raw_status == "closed":
                section = "Closed"
            if raw_status == "elapsed-rejected":
                section = "Elapsed-Rejected"
            flags = list(record.flags) + (list(ranked.review_flags) if ranked else [])
            flags += [str(v[h]) for h in ("Assistant extraction flags", "Assistant fetch issue", "Assistant discovery notes", "Assistant assessment issue", "Assistant assessment uncertainties") if v.get(h)]
            # Keep prior extraction warnings visible on repeat runs, but recompute
            # engine warnings and queue reasons from current inputs.
            v.update({"Assistant section": section, "Assistant priority": ranked.priority if ranked else None,
                      "Assistant overall fit": ranked.overall_fit if ranked else "unknown",
                      "Assistant location eligibility": ranked.location_eligibility if ranked else "unknown",
                      "Assistant priority reasons": "\n".join(ranked.reasons) if ranked else "Not ranked until the review issue is resolved.",
                      "Assistant review flags": "\n".join(dict.fromkeys(flags)), "Assistant evaluated on": as_of,
                      "Assistant policy version": ranked.rubric_version if ranked else "1.0"})
            known_criteria = sum(getattr(record.job.assessment, key) != "unknown" for key in
                                 ("leadership_fit", "technical_fit", "career_alignment", "opportunity_quality")) if record.job else 0
            v["Assistant assessment status"] = ("Complete" if known_criteria == 4 else "Partial" if known_criteria else "Needs assessment")
            if raw_status == "elapsed-rejected":
                v["Assistant priority reasons"] = "Application aged at least one calendar month: Elapsed-Rejected is a tracking assumption, not a confirmed employer rejection."
        source_order = {id(record): i for i, record in enumerate(records)}
        section_order = {"Interviewing": -1, "Applied": 0, "Outstanding": 1, "Review": 2, "Closed": 3, "Rejected": 4, "Elapsed-Rejected": 5}
        output_records.sort(key=lambda r: (section_order[r.values["Assistant section"]],
                            r.values["Assistant priority"] if r.values["Assistant section"] == "Outstanding" else source_order[id(r)]))
        # A numeric recommendation requires an actual contextual assessment.
        # Missing evidence must not masquerade as a low fit score or priority.
        apply_order = 0
        for record in output_records:
            status = str(_value(record, "status") or "").strip().casefold()
            if status in ("", "not applied", "outstanding"):
                if record.values["Assistant section"] == "Outstanding" and record.values["Assistant assessment status"] != "Needs assessment":
                    apply_order += 1
                    record.values["Apply order"] = apply_order
                    record.values["Assistant priority"] = apply_order
                else:
                    record.values["Apply order"] = "Needs assessment" if record.values["Assistant assessment status"] == "Needs assessment" else "Needs review"
                    record.values["Assistant priority"] = None
                    record.values["Assistant priority reasons"] = "No numeric apply recommendation until assessment/review is resolved.\n" + record.values["Assistant priority reasons"]
            else:
                record.values["Apply order"] = "-"
        def display_order(record):
            section = record.values["Assistant section"]
            order = record.values["Apply order"]
            if section == "Interviewing":
                return (-1, source_order[id(record)])
            if section == "Applied":
                return (0, source_order[id(record)])
            if isinstance(order, int):
                return (1, order)
            return ({"Closed": 3, "Rejected": 4, "Elapsed-Rejected": 5}.get(section, 2), section_order[section], source_order[id(record)])
        output_records.sort(key=display_order)
        if QUEUE_SHEET in wb.sheetnames:
            index = 1
            while f"Queue archive {index}" in wb.sheetnames:
                index += 1
            wb[QUEUE_SHEET].title = f"Queue archive {index}"
        target = wb.create_sheet(QUEUE_SHEET, 0)
        all_headers = list(headers)
        for label in MANUAL_COLUMNS + GENERATED_COLUMNS:
            if label not in all_headers:
                all_headers.append(label)
        target.append(all_headers)
        source_columns = {h: i for i, h in enumerate(headers, 1)}
        for new_row, record in enumerate(output_records, 2):
            for col, label in enumerate(all_headers, 1):
                cell = target.cell(new_row, col)
                if record.row is not None and label in source_columns and label not in GENERATED_COLUMNS:
                    original = ws.cell(record.row, source_columns[label])
                    cell.value = original.value
                    cell.data_type = original.data_type
                    if original.data_type == "f":
                        cell.value = Translator(original.value, origin=original.coordinate).translate_formula(cell.coordinate)
                    cell._style = copy(original._style)
                    cell.hyperlink = copy(original.hyperlink)
                    cell.comment = copy(original.comment)
                    cell.protection = copy(original.protection)
                    if label in ALIASES["status"] + ALIASES["company"] + ALIASES["title"] and record.values.get(label) != original.value:
                        _write_text(cell, record.values.get(label))
                else:
                    value = record.values.get(label)
                    if isinstance(value, str) and len(value) > 32767:
                        raise ValueError(f"Cell exceeds Excel's text limit: {label}; original workbook is unchanged")
                    _write_text(cell, value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd"
                if label in ALIASES["url"] and cell.value and not cell.hyperlink and str(cell.value).startswith(("http://", "https://")):
                    cell.hyperlink = str(cell.value)
            target.row_dimensions[new_row].height = 48
        colors = {"Interviewing": "DDEBF7", "Applied": "E2F0D9", "Outstanding": "E7EFFA", "Rejected": "FCE4D6", "Review": "FFF2CC", "Closed": "D9D9D9", "Elapsed-Rejected": "FCE4D6"}
        section_col = all_headers.index("Assistant section") + 1
        for row in target.iter_rows(min_row=2):
            row[section_col-1].fill = PatternFill("solid", fgColor=colors[row[section_col-1].value])
        for cell in target[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244062")
            cell.alignment = Alignment(wrap_text=True)
        for i, label in enumerate(all_headers, 1):
            target.column_dimensions[get_column_letter(i)].width = 42 if any(word in label.lower() for word in ("reason", "evidence", "description", "flags", "gap")) else 24
        target.freeze_panes = "D2"
        target.auto_filter.ref = target.dimensions
        target.sheet_properties.pageSetUpPr.fitToPage = True
        if compact:
            if any(c.data_type == "f" and "!" in str(c.value) for row in target for c in row):
                raise ValueError("Single-tab export cannot preserve cross-sheet formulas; use --keep-history-tabs")
            for old_sheet in list(wb.worksheets):
                if old_sheet is not target:
                    wb.remove(old_sheet)
        wb.active = 0
        output.parent.mkdir(parents=True, exist_ok=True)
        # Exclusive creation closes the race between existence check and save.
        with output.open("xb") as stream:
            try:
                wb.save(stream)
            except Exception:
                stream.close()
                output.unlink(missing_ok=True)
                raise
        review_count = sum(r.values["Assistant section"] == "Review" for r in records)
        return UpdateSummary(output, source_sheet_name, imported, added, skipped, failures, len(result.outstanding), review_count)
    finally:
        wb.close()
