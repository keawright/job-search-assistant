from copy import copy
from datetime import date
import hashlib
import json

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
import pytest

from job_search_assistant.cli import main
from job_search_assistant.ingestion import FetchError, Page
from job_search_assistant.workbook import update_workbook

TODAY = date(2026, 10, 15)
HEADERS = ["Apply order", "Company", "Position title", "Job posting (URL)", "Application status", "Applied Date", "Outcome Reason", "Personal notes", "Deadline"]


@pytest.fixture
def workbook(tmp_path):
    source = tmp_path / "input.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Updated priorities"
    ws.append(["Historical source notes"])
    ws.append(HEADERS)
    ws.append([1, "Example", "EM", "https://example.com/job/1", "Not applied", None, None, "Keep this note"])
    ws.append([None, "Applied Co", "Director", "https://example.com/job/2", "Applied", "Oct-9-2026", None, "Referral"])
    ws.append([None, "Rejected Co", "PM", "https://example.com/job/3", "Rejected", None, "User ruled out", "Keep"])
    ws.append([None, "Recruiter", "SEM", None, "Applied", "Oct-9-2026", None, "Recruiter discussion"])
    ws["H3"].comment = Comment("Manual annotation", "Candidate")
    ws["D3"].hyperlink = "https://example.com/job/1"
    wb.create_sheet("Source archive")["A1"] = "Untouched"
    wb.save(source)
    wb.close()
    return source


def rows(path):
    wb = load_workbook(path)
    ws = wb["Assistant queue"]
    headers = [c.value for c in ws[1]]
    records = [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True)]
    return wb, records


def fetched(url):
    data = {"@type": "JobPosting", "title": "Engineering Manager", "hiringOrganization": {"name": "Synthetic Labs"},
            "description": "<p>Lead SaaS teams.</p>", "datePosted": "2026-10-01", "jobLocationType": "TELECOMMUTE",
            "applicantLocationRequirements": {"@type": "Country", "name": "Canada"}}
    return Page(url, '<script type="application/ld+json">' + json.dumps(data) + '</script>')


def test_update_preserves_original_tabs_tracking_and_missing_url_applied(workbook, tmp_path):
    before = hashlib.sha256(workbook.read_bytes()).hexdigest()
    output = tmp_path / "updated.xlsx"
    summary = update_workbook(workbook, output, ["https://example.com/new"], as_of=TODAY, fetcher=fetched)
    assert summary.imported == 4 and summary.added == 1
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == before
    wb, records = rows(output)
    try:
        assert wb["Source archive"]["A1"].value == "Untouched"
        assert wb["Updated priorities"]["H3"].value == "Keep this note"
        assert wb["Updated priorities"]["H3"].comment.text == "Manual annotation"
        assert [r["Company"] for r in records[:2]] == ["Applied Co", "Recruiter"]
        applied = records[0]
        assert applied["Application status"] == "Applied" and applied["Applied Date"] == "Oct-9-2026"
        assert records[1]["Assistant section"] == "Applied"
        new = next(r for r in records if r["Company"] == "Synthetic Labs")
        assert new["Assistant overall fit"] == "unknown"
        assert new["Assistant location eligibility"] == "eligible"
        snapshot = json.loads(__import__('pathlib').Path(new["Assistant snapshot"]).read_text(encoding="utf-8"))
        assert snapshot["extracted"]["description"] == "Lead SaaS teams."
        assert "<script" in snapshot["html"]
    finally:
        wb.close()


def test_rerun_uses_latest_queue_preserves_edits_and_skips_tracking_alias(workbook, tmp_path):
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    update_workbook(workbook, first, ["https://example.com/new"], as_of=TODAY, fetcher=fetched)
    wb = load_workbook(first)
    ws = wb["Assistant queue"]
    labels = {cell.value: cell.column for cell in ws[1]}
    for row in range(2, ws.max_row+1):
        if ws.cell(row, labels["Company"]).value == "Synthetic Labs":
            ws.cell(row, labels["Application status"], "Applied")
            ws.cell(row, labels["Personal notes"], "My new note")
    wb.save(first)
    wb.close()
    result = update_workbook(first, second, ["https://example.com/new?utm_source=newsletter#top"], as_of=TODAY, offline=True)
    assert result.added == 0 and result.skipped_urls == 1
    wb, records = rows(second)
    try:
        assert "Queue archive 1" in wb.sheetnames
        new = next(r for r in records if r["Company"] == "Synthetic Labs")
        assert new["Assistant section"] == "Applied" and new["Personal notes"] == "My new note"
        assert "Source authority" in new["Assistant review flags"]
    finally:
        wb.close()


def test_existing_output_and_input_never_overwritten(workbook, tmp_path):
    with pytest.raises(ValueError, match="never overwritten"):
        update_workbook(workbook, workbook, [], as_of=TODAY)
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"important")
    with pytest.raises(ValueError):
        update_workbook(workbook, output, [], as_of=TODAY)
    assert output.read_bytes() == b"important"


def test_offline_does_not_fetch_and_batch_duplicates_skip(workbook, tmp_path):
    def forbidden(url):
        raise AssertionError("Network used in offline mode")
    output = tmp_path / "offline.xlsx"
    result = update_workbook(workbook, output, ["https://example.com/new", "https://example.com/new?utm_source=x", "https://example.com/job/1"],
                             as_of=TODAY, offline=True, refresh=True, fetcher=forbidden)
    assert result.added == 1 and result.skipped_urls == 2


def test_failed_fetch_remains_visible_and_not_closed(workbook, tmp_path):
    def failed(url):
        raise FetchError("HTTP 403")
    output = tmp_path / "failed.xlsx"
    result = update_workbook(workbook, output, ["https://example.com/blocked"], as_of=TODAY, fetcher=failed)
    assert result.fetch_failures == 2
    wb, records = rows(output)
    try:
        new = next(r for r in records if r["Job posting (URL)"] == "https://example.com/blocked")
        assert "HTTP 403" in new["Assistant review flags"]
        assert new["Assistant overall fit"] == "unknown"
    finally:
        wb.close()


def test_unrecognized_status_and_duplicates_retained_for_review(workbook, tmp_path):
    wb = load_workbook(workbook)
    ws = wb["Updated priorities"]
    ws.append([None, "Duplicate", "PM", "https://example.com/job/1", "Not applied"])
    ws.append([None, "Custom status", "PM", "https://example.com/custom", "Offer pending"])
    wb.save(workbook)
    wb.close()
    output = tmp_path / "review.xlsx"
    result = update_workbook(workbook, output, [], as_of=TODAY, offline=True)
    assert result.review == 2
    wb, records = rows(output)
    try:
        assert len(records) == 6
        custom = next(r for r in records if r["Company"] == "Custom status")
        assert custom["Application status"] == "Offer pending"
        assert custom["Assistant section"] == "Review"
    finally:
        wb.close()


def test_page_text_cannot_inject_excel_formulas(workbook, tmp_path):
    def malicious(url):
        return Page(url, '<script type="application/ld+json">' + json.dumps({"@type": "JobPosting", "title": '=HYPERLINK("http://bad", "click")', "description": "=1+1"}) + '</script>')
    output = tmp_path / "safe.xlsx"
    update_workbook(workbook, output, ["https://example.com/new"], as_of=TODAY, fetcher=malicious)
    wb, records = rows(output)
    try:
        ws = wb["Assistant queue"]
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert cell.data_type == "s"
    finally:
        wb.close()


def test_editable_assessment_and_manual_priority_affect_next_run(workbook, tmp_path):
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    update_workbook(workbook, first, [], as_of=TODAY, offline=True)
    wb = load_workbook(first)
    ws = wb["Assistant queue"]
    labels = {c.value: c.column for c in ws[1]}
    row = next(i for i in range(2, ws.max_row+1) if ws.cell(i, labels["Company"]).value == "Example")
    for label in ("Leadership fit", "Technical fit", "Career alignment", "Opportunity quality"):
        ws.cell(row, labels[label], "good")
    ws.cell(row, labels["Assessment evidence"], "Comparable SaaS team leadership and delivery; scope is a reasonable progression.")
    ws.cell(row, labels["Manual priority"], 1)
    wb.save(first)
    wb.close()
    update_workbook(first, second, [], as_of=TODAY, offline=True)
    wb, records = rows(second)
    try:
        assessed = next(r for r in records if r["Company"] == "Example")
        assert assessed["Assistant overall fit"] == "good"
        assert assessed["Assistant priority"] == 1
        assert assessed["Personal notes"] == "Keep this note"
    finally:
        wb.close()


def test_cli_url_file_and_preferences(workbook, tmp_path, capsys):
    urls, preferences, output = tmp_path / "urls.txt", tmp_path / "preferences.json", tmp_path / "cli.xlsx"
    urls.write_text("# Two new postings\nhttps://example.com/new1\nhttps://example.com/new2\n", encoding="utf-8")
    preferences.write_text(json.dumps({"country": "Canada", "prefer_leadership": True}), encoding="utf-8")
    assert main(["update", "--workbook", str(workbook), "--output", str(output), "--urls-file", str(urls), "--preferences", str(preferences), "--offline", "--as-of", TODAY.isoformat()]) == 0
    assert output.exists() and "added 2" in capsys.readouterr().out
    assert main(["update", "--workbook", str(workbook), "--output", str(output)]) == 2


def test_redirect_alias_duplicate_not_added(workbook, tmp_path):
    def alias(url):
        return Page("https://example.com/job/1", '<main>Existing posting</main>')
    result = update_workbook(workbook, tmp_path / "alias.xlsx", ["https://example.com/alias"], as_of=TODAY, fetcher=alias)
    assert result.added == 0 and result.skipped_urls == 1
