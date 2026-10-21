from datetime import date

from openpyxl import Workbook, load_workbook

from job_search_assistant.workbook import update_workbook


def read(path):
    wb = load_workbook(path)
    values = list(wb["Assistant queue"].values)
    result = [dict(zip(values[0], row)) for row in values[1:]]
    wb.close()
    return result


def test_assessments_drive_apply_order_and_reassessment_changes_it(tmp_path):
    source, first, second = [tmp_path / name for name in ("source.xlsx", "first.xlsx", "second.xlsx")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Assistant queue"
    criteria = ["Leadership fit", "Technical fit", "Career alignment", "Opportunity quality"]
    ws.append(["Company", "Position title", "Job posting", "Application status", *criteria, "Assessment evidence", "Why this fits"])
    for name, rating in [("A stretch", "stretch"), ("Z strong", "strong"), ("Unassessed", None)]:
        ws.append([name, "EM", f"https://example.com/{name.replace(' ', '-')}", "Not applied", *([rating] * 4), "Role responsibilities compared with candidate accomplishments." if rating else None, "Legacy commentary alone is not a criterion assessment"])
    wb.save(source)
    wb.close()
    update_workbook(source, first, [], as_of=date(2026, 10, 20), offline=True)
    rows = read(first)
    assert [r["Company"] for r in rows] == ["Z strong", "A stretch", "Unassessed"]
    assert [r["Apply order"] for r in rows] == [1, 2, "Needs assessment"]
    assert rows[-1]["Assistant priority"] is None
    assert rows[-1]["Assistant assessment status"] == "Needs assessment"
    assert rows[-1]["Why this fits"] == "Legacy commentary alone is not a criterion assessment"
    wb = load_workbook(first)
    ws = wb["Assistant queue"]
    cols = {cell.value: cell.column for cell in ws[1]}
    for row in range(2, ws.max_row + 1):
        company = ws.cell(row, cols["Company"]).value
        if company in ("A stretch", "Z strong"):
            for label in criteria:
                ws.cell(row, cols[label], "strong" if company == "A stretch" else "stretch")
    wb.save(first)
    wb.close()
    update_workbook(first, second, [], as_of=date(2026, 10, 20), offline=True)
    rows = read(second)
    assert [r["Company"] for r in rows] == ["A stretch", "Z strong", "Unassessed"]
    assert [r["Apply order"] for r in rows] == [1, 2, "Needs assessment"]
