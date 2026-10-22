from copy import deepcopy
from dataclasses import replace
from datetime import date
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
import pytest

from job_search_assistant.assessment import (
    AssessmentError, AssessmentService, CRITERIA, CodexAssessor, OpenAIAssessor,
    assessment_key, job_evidence, load_service, validate_assessment,
)
from job_search_assistant.models import CandidatePreferences, Job
from job_search_assistant.workbook import update_workbook

TODAY = date(2026, 10, 20)
RESUME = 'Led two distributed cloud engineering teams. Delivered reliable infrastructure and developed senior engineers. Program leadership and operational improvements.'
DESCRIPTION = 'Lead distributed engineering teams and develop senior engineers. Own reliable cloud infrastructure and operational improvement programs.'


@pytest.fixture
def job():
    return Job('test-1', 'Example', 'Engineering Manager', 'https://example.com/job', description=DESCRIPTION)


@pytest.fixture
def payload():
    return {**{key: {'rating':'good', 'resume_quote':'Led two distributed cloud engineering teams.',
                      'job_quote':'Lead distributed engineering teams and develop senior engineers.',
                      'rationale':'The demonstrated distributed-team leadership transfers to the stated role.'} for key in CRITERIA},
            'strengths':['Distributed leadership'], 'gaps':[], 'regular_hands_on_coding':None, 'coding_quote':'',
            'summary':'Relevant distributed cloud leadership.', 'uncertainties':['Confirm specific scope and current vacancy.']}


def test_quotes_and_types_are_validated(job, payload):
    assert validate_assessment(payload, RESUME, job_evidence(job)).leadership_fit == 'good'
    payload['technical_fit']['resume_quote'] = 'Invented production ML achievement'
    with pytest.raises(AssessmentError, match='unsupported'):
        validate_assessment(payload, RESUME, job_evidence(job))


@pytest.mark.parametrize('mutation', ['coding', 'gap', 'unknown_field', 'rating', 'list'])
def test_invalid_assessments_never_rank(job, payload, mutation):
    if mutation == 'coding': payload['regular_hands_on_coding'] = True
    elif mutation == 'gap': payload['gaps'] = [{'description':'Gap', 'severity':'major', 'resume_quote':'Invented evidence', 'job_quote':'Invented evidence', 'rationale':'Unsupported'}]
    elif mutation == 'unknown_field': payload['extra'] = 'ignored instructions'
    elif mutation == 'rating': payload['leadership_fit']['rating'] = 'perfect'
    else: payload['strengths'] = 'not a list'
    with pytest.raises((AssessmentError, ValueError)):
        validate_assessment(payload, RESUME, job_evidence(job))


def test_saved_assessment_reused_offline_and_invalidated_by_inputs(tmp_path, job, payload):
    candidate = CandidatePreferences()
    service = AssessmentService(tmp_path, RESUME, allow_network=False)
    service.save(job, candidate, payload, source='Conversational review', assessed_on=TODAY)
    rating, entry = service.resolve(job, candidate, TODAY, offline=True)
    assert rating.leadership_fit == 'good' and service.reused == 1
    assert entry['source'] == 'Conversational review'
    changed = replace(job, description=DESCRIPTION + ' Requires new specialist work.')
    with pytest.raises(AssessmentError, match='No matching'):
        service.resolve(changed, candidate, TODAY)
    assert assessment_key(job, candidate, RESUME) != assessment_key(job, candidate, RESUME + ' More experience.')
    assert assessment_key(job, candidate, RESUME) != assessment_key(job, replace(candidate, prefer_leadership=False), RESUME)


def test_cache_tampering_revalidated(tmp_path, job, payload):
    service = AssessmentService(tmp_path, RESUME)
    candidate = CandidatePreferences()
    service.save(job,candidate,payload,source='Review',assessed_on=TODAY)
    path = next((tmp_path/'assessments').glob('*.json'))
    data = json.loads(path.read_text())
    data['assessment']['technical_fit']['job_quote'] = 'Invented job requirement'
    path.write_text(json.dumps(data))
    with pytest.raises(AssessmentError): service.resolve(job,candidate,TODAY)


def test_api_contract_refusals_and_incomplete(job, payload):
    calls = []
    def transport(request, key):
        calls.append(request)
        assert key == 'synthetic-key'
        return {'status':'completed','output':[{'type':'message','content':[{'type':'output_text','text':json.dumps(payload)}]}]}
    provider = OpenAIAssessor('synthetic-key',transport=transport)
    assert provider.assess(RESUME,job_evidence(job),CandidatePreferences()) == payload
    assert calls[0]['store'] is False
    assert calls[0]['text']['format']['strict'] is True
    assert 'untrusted data' in calls[0]['instructions']
    for response in [{'status':'incomplete'}, {'status':'completed','output':[{'type':'message','content':[{'type':'refusal','refusal':'No'}]}]}]:
        provider = OpenAIAssessor('synthetic-key',transport=lambda *args: response)
        with pytest.raises(AssessmentError): provider.assess(RESUME,job_evidence(job),CandidatePreferences())


def test_codex_uses_stdin_no_shell_schema_readonly_and_existing_login(tmp_path, monkeypatch, job, payload):
    monkeypatch.setattr('job_search_assistant.assessment.codex_command',lambda:['codex'])
    def runner(command, **kwargs):
        assert command[0:2] == ['codex','exec']
        assert '--ephemeral' in command and '--ignore-user-config' in command
        assert command[command.index('--sandbox')+1] == 'read-only'
        assert 'features.shell_tool=false' in command
        assert 'features.multi_agent=false' in command
        assert RESUME in kwargs['input']
        assert kwargs['shell'] is False
        Path(command[command.index('--output-last-message')+1]).write_text(json.dumps(payload),encoding='utf-8')
        return SimpleNamespace(returncode=0)
    assessor = CodexAssessor(tmp_path,runner=runner)
    assert assessor.assess(RESUME,job_evidence(job),CandidatePreferences()) == payload
    assert list((tmp_path/'codex-runs').iterdir()) == []


def test_call_limit_and_offline_prevent_provider_requests(tmp_path, job, payload):
    class Fake:
        name, model = 'Synthetic', 'test'
        calls = 0
        def assess(self, *args): self.calls += 1; return deepcopy(payload)
    provider = Fake()
    service = AssessmentService(tmp_path,RESUME,provider,max_calls=1)
    with pytest.raises(AssessmentError): service.resolve(job,CandidatePreferences(),TODAY,offline=True)
    assert provider.calls == 0
    service.resolve(job,CandidatePreferences(),TODAY)
    assert provider.calls == 1
    with pytest.raises(AssessmentError,match='limit'): service.resolve(replace(job,url='https://example.com/new'),CandidatePreferences(),TODAY)


def test_codex_failures_do_not_leak_raw_process_output(tmp_path, monkeypatch, job):
    monkeypatch.setattr('job_search_assistant.assessment.codex_command',lambda:['codex'])
    assessor = CodexAssessor(tmp_path,runner=lambda *args,**kwargs:SimpleNamespace(returncode=1,stderr='private diagnostic text'))
    with pytest.raises(AssessmentError) as failure:
        assessor.assess(RESUME,job_evidence(job),CandidatePreferences())
    assert 'private diagnostic' not in str(failure.value)


def test_failed_ai_leaves_no_numeric_recommendation(tmp_path):
    source,output=tmp_path/'source.xlsx',tmp_path/'output.xlsx'
    wb=Workbook();ws=wb.active;ws.title='Assistant queue'
    ws.append(['Company','Position title','Job posting','Application status','Assistant description'])
    ws.append(['Example','EM','https://example.com/1','Not applied',DESCRIPTION])
    wb.save(source);wb.close()
    class Failed:
        name,model='Synthetic','test'
        def assess(self,*args): raise AssessmentError('Codex assessment failed; check login')
    service=AssessmentService(tmp_path/'private',RESUME,Failed())
    from job_search_assistant.ingestion import FetchError
    def fetch(url): raise FetchError('Blocked')
    update_workbook(source,output,[],as_of=TODAY,fetcher=fetch,assessment_service=service)
    wb=load_workbook(output);data=list(wb['Assistant queue'].values);row=dict(zip(data[0],data[1]));wb.close()
    assert row['Application status']=='Not applied'
    assert row['Apply order']=='Needs assessment'
    assert row['Assistant priority'] is None
    assert 'check login' in row['Assistant assessment issue']


def test_similar_jobs_removed_from_assessment_input(job):
    context = job_evidence(replace(job,description=DESCRIPTION+' Similar jobs Unrelated role requiring machine learning'))
    assert 'Unrelated' not in context['description']


def test_default_codex_config_and_offline_loading(tmp_path):
    private = tmp_path/'.job-search-assistant-private'; private.mkdir()
    (private/'resume.txt').write_text(RESUME)
    (private/'assessment-config.json').write_text(json.dumps({'resume_path':'.job-search-assistant-private/resume.txt','provider':'codex'}))
    assert isinstance(load_service(tmp_path).provider,CodexAssessor)
    assert load_service(tmp_path,offline=True).provider is None
    assert load_service(tmp_path,no_ai=True).allow_network is False


def test_workbook_uses_ai_ratings_without_changing_manual_statuses_or_fields(tmp_path, payload):
    source, output = tmp_path/'input.xlsx',tmp_path/'output.xlsx'
    wb = Workbook(); ws=wb.active; ws.title='Assistant queue'
    ws.append(['Company','Position title','Job posting','Application status','Assistant description','Leadership fit','Assessment evidence','My notes'])
    for name,status in [('Auto','Not applied'),('Manual','Not applied'),('Applied','Applied'),('Closed','Closed'),('Rejected','Rejected'),('Elapsed','Elapsed-Rejected'),('Interview','Interviewing')]:
        ws.append([name,'EM',f'https://example.com/{name}',status,DESCRIPTION,'strong' if name=='Manual' else None,'Manual comparison evidence' if name=='Manual' else None,'Keep'])
    wb.save(source);wb.close()
    class Fake:
        name,model='Synthetic','test'
        calls=0
        def assess(self,*args): self.calls+=1; return deepcopy(payload)
    provider=Fake();service=AssessmentService(tmp_path/'private',RESUME,provider)
    def fetch(url):
        from job_search_assistant.ingestion import FetchError
        raise FetchError('Blocked, preserve the description')
    update_workbook(source,output,[],as_of=TODAY,fetcher=fetch,assessment_service=service)
    assert provider.calls==1
    wb=load_workbook(output); data=list(wb['Assistant queue'].values); rows={v[0]:dict(zip(data[0],v)) for v in data[1:]}
    assert rows['Auto']['Assistant overall fit']=='good'
    assert rows['Auto']['Leadership fit'] is None
    assert isinstance(rows['Auto']['Apply order'],int)
    assert rows['Manual']['Assistant overall fit']=='strong'
    assert rows['Manual']['Assessment evidence']=='Manual comparison evidence'
    for name,status in [('Applied','Applied'),('Closed','Closed'),('Rejected','Rejected'),('Elapsed','Elapsed-Rejected'),('Interview','Interviewing')]:
        assert rows[name]['Application status']==status
    assert all(row['My notes']=='Keep' for row in rows.values())
    wb.close()


def test_selected_rows_force_only_selected_assessments(tmp_path, payload):
    source, output = tmp_path/'input.xlsx', tmp_path/'output.xlsx'
    wb=Workbook(); ws=wb.active; ws.title='Assistant queue'
    ws.append(['Company','Position title','Job posting','Application status','Assistant description'])
    for i in range(3):
        ws.append(['Example','Engineering Manager',f'https://example.com/{i}','Not applied',DESCRIPTION])
    wb.save(source); wb.close()
    from job_search_assistant.workbook import _read_records, _make_job
    class Fake:
        name, model = 'Synthetic', 'test'
        def __init__(self): self.urls=[]
        def assess(self, resume, context, candidate):
            self.urls.append(context['url']); return deepcopy(payload)
    provider=Fake();service=AssessmentService(tmp_path/'private',RESUME,provider)
    wb=load_workbook(source); _,_,records=_read_records(wb,None)
    # Selected and unselected rows both have valid cached assessments.
    for record in records[:2]:
        service.save(_make_job(record,wb.epoch),CandidatePreferences(),deepcopy(payload),source='Previous',assessed_on=TODAY)
    wb.close()
    from job_search_assistant.ingestion import Page
    fetched=[]
    def fetch(url):
        fetched.append(url)
        from job_search_assistant.ingestion import FetchError
        raise FetchError("Blocked; keep saved description")
    update_workbook(source,output,[],as_of=TODAY,fetcher=fetch,assessment_service=service,reassess_rows={2})
    assert provider.urls == ['https://example.com/0']
    assert fetched == ['https://example.com/0']
    assert service.reused == 1
    wb=load_workbook(output);data=list(wb['Assistant queue'].values)
    rows=[dict(zip(data[0],r)) for r in data[1:]];wb.close()
    assert sum(isinstance(r['Apply order'],int) for r in rows)==2
    assert any('disabled' in str(r['Assistant assessment issue']) for r in rows)


@pytest.mark.parametrize('selection', ['1','3-2','abc','2-3-4','2,','1048577'])
def test_invalid_reassessment_rows(selection):
    import argparse
    from job_search_assistant.cli import _row_selection
    with pytest.raises(argparse.ArgumentTypeError): _row_selection(selection)


def test_reassessment_row_ranges():
    from job_search_assistant.cli import _row_selection
    assert _row_selection('98,101,103-105,104') == {98,101,103,104,105}


@pytest.mark.parametrize('reverse', [False,True])
def test_reassess_needed_follows_saved_status_after_rows_move(tmp_path,payload,reverse):
    source,output=tmp_path/'input.xlsx',tmp_path/'output.xlsx'
    wb=Workbook();ws=wb.active;ws.title='Assistant queue'
    ws.append(['Company','Position title','Job posting','Application status','Assistant description','Apply order'])
    items=[('assess','Not applied',' Needs ASSESSMENT '),('review','Not applied','Needs review'),
           ('typo','Not applied','Needs Assessement'),('ranked','Not applied',1),
           ('closed','Closed','Needs assessment'),('interview','Interviewing','Needs review')]
    for name,status,order in reversed(items) if reverse else items:
        ws.append(['Example','Engineering Manager',f'https://example.com/{name}',status,DESCRIPTION,order])
    wb.save(source);wb.close()
    class Fake:
        name,model='Synthetic','test'
        def __init__(self): self.urls=[]
        def assess(self,resume,context,candidate):
            self.urls.append(context['url']);return deepcopy(payload)
    provider=Fake();service=AssessmentService(tmp_path/'private',RESUME,provider)
    from job_search_assistant.ingestion import FetchError
    def fetch(url): raise FetchError('Blocked; retain saved description')
    update_workbook(source,output,[],as_of=TODAY,fetcher=fetch,assessment_service=service,reassess_needed=True)
    assert set(provider.urls)=={f'https://example.com/{name}' for name in ('assess','review','typo')}
    assert len(provider.urls)==3


@pytest.mark.parametrize("limit", [1, 2])
def test_new_jobs_assessed_first_and_existing_cache_reused(tmp_path, payload, limit):
    source, output = tmp_path/'input.xlsx', tmp_path/'output.xlsx'
    wb=Workbook(); ws=wb.active; ws.title='Assistant queue'
    ws.append(['Company','Position title','Job posting','Application status','Assistant description','My notes'])
    for name in ['Missing', 'Cached']:
        ws.append([name,'Engineering Manager',f'https://example.com/{name}','Not applied',DESCRIPTION,'Keep'])
    wb.save(source)
    class Fake:
        name, model = 'Synthetic', 'test'
        def __init__(self): self.urls=[]
        def assess(self, resume, context, candidate):
            self.urls.append(context['url']); return deepcopy(payload)
    provider=Fake()
    service=AssessmentService(tmp_path/'private',RESUME,provider,max_calls=limit)
    from job_search_assistant.workbook import _read_records, _make_job
    _,_,records=_read_records(wb,None)
    service.save(_make_job(records[1],wb.epoch),CandidatePreferences(),deepcopy(payload),source='Previous',assessed_on=TODAY)
    wb.close()
    from job_search_assistant.ingestion import Page, FetchError
    def fetch(url):
        if url != 'https://example.com/new':
            raise FetchError('Blocked; keep saved description')
        data={'@type':'JobPosting','title':'Engineering Manager',
              'hiringOrganization':{'name':'New'},'description':DESCRIPTION}
        return Page(url, '<script type="application/ld+json">'+json.dumps(data)+'</script>')
    messages=[]
    result=update_workbook(source,output,['https://example.com/new'],as_of=TODAY,
                           fetcher=fetch,assessment_service=service,progress=messages.append)
    assert provider.urls == ['https://example.com/new'] + (['https://example.com/Missing'] if limit == 2 else [])
    assert service.reused == 1
    assert service.created == limit
    assert sum(result.assessment_issues.values()) == (1 if limit == 1 else 0)
    assert any('New job assessment: New' in message for message in messages)
    assert any('saved result reused' in message for message in messages)
    assert any('call limit' in message for message in messages) == (limit == 1)
    wb=load_workbook(output); data=list(wb['Assistant queue'].values)
    rows={r[0]:dict(zip(data[0],r)) for r in data[1:]}; wb.close()
    assert rows['New']['Assistant assessment status'] == 'Complete'
    assert rows['Cached']['Assistant assessment status'] == 'Complete'
    assert rows['Missing']['My notes'] == 'Keep'
    assert rows['Missing']['Application status'] == 'Not applied'
    assert rows['Missing']['Assistant assessment status'] == ('Complete' if limit == 2 else 'Needs assessment')


def test_default_initial_assessments_exceed_twenty_and_refresh_budget_is_separate(tmp_path, job, payload):
    class Fake:
        name, model = 'Synthetic', 'test'
        def assess(self, *args): return deepcopy(payload)
    candidate=CandidatePreferences()
    service=AssessmentService(tmp_path, RESUME, Fake(), max_refresh_calls=1)
    for i in range(23):
        service.resolve(replace(job,id=f'new-{i}',url=f'https://example.com/{i}'),candidate,TODAY)
    assert service.created == 23
    # A new run finds prior successful assessments by stable ID, even with changed URLs.
    service=AssessmentService(tmp_path, RESUME, Fake(), max_refresh_calls=1)
    first=replace(job,id='new-0',url='https://example.com/changed0')
    service.resolve(first,candidate,TODAY)
    assert service.refresh_calls == 1
    with pytest.raises(AssessmentError, match='refresh limit'):
        service.resolve(replace(job,id='new-1',url='https://example.com/changed1'),candidate,TODAY)
    # Cache reuse and initial assessments continue after refresh exhaustion.
    service.resolve(first,candidate,TODAY)
    service.resolve(replace(job,id='another-new'),candidate,TODAY)
    assert service.reused == 1
    assert service.created == 2


def test_legacy_config_does_not_cap_initial_assessments(tmp_path):
    private=tmp_path/'.job-search-assistant-private'; private.mkdir()
    (private/'resume.txt').write_text(RESUME)
    config={'resume_path':'.job-search-assistant-private/resume.txt','max_calls':1}
    path=private/'assessment-config.json';path.write_text(json.dumps(config))
    service=load_service(tmp_path)
    assert service.max_calls is None
    assert service.max_refresh_calls == 1
    assert load_service(tmp_path,max_calls=3).max_calls == 3
    config['max_refresh_calls']=2;path.write_text(json.dumps(config))
    assert load_service(tmp_path).max_refresh_calls == 2
