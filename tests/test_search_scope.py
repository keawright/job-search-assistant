from types import SimpleNamespace
from openpyxl import Workbook, load_workbook
import pytest

from job_search_assistant.searching.eligibility import title_gate, location_gate
from job_search_assistant.searching.normalization import Posting
from job_search_assistant.searching.reporting import export


@pytest.mark.parametrize('title',[
    'Director of Engineering','Associate Director, Technology','Associate Director',
    'Organizational Change Manager','Organizational Change Management Lead',
    'Delivery Manager','Technical Delivery Manager','Engineering Manager',
    'Senior Engineering Manager','Senior Project Manager','Senior Program Manager',
    'Director of Technical Program Management','Manager, Software Engineering',
])
def test_requested_leadership_roles_are_retained(title):
    assert title_gate(title)


@pytest.mark.parametrize('title',[
    'Senior Software Engineer, Backend (Identity International)',
    'Staff Software Engineer','Product Security Engineer II - IAM',
    'Senior Security Engineer','Lead Software Engineer','Principal Data Engineer',
    'Software Engineer II, Frontend','Senior Site Reliability Engineer',
])
def test_individual_contributor_roles_are_outside_scope(title):
    assert not title_gate(title)


@pytest.mark.parametrize('location',[
    'Remote Spain','Remote Poland','Remote US','Remote USA','Remote UK','Warsaw, Poland',
    'Madrid, Spain','Remote Germany','Hybrid - Toronto, Ontario','Office Based - Taipei, Taiwan',
])
def test_foreign_and_distant_hybrid_labels_do_not_become_canadian_eligibility(location):
    assert location_gate('Our company also has employees in Canada.',location)[0]=='ineligible'


@pytest.mark.parametrize('location',['Remote Canada','Remote US / Canada','Remote Canada / Spain'])
def test_explicit_canadian_multi_country_options_remain_allowed(location):
    assert location_gate('Lead distributed engineering teams.',location)[0]=='eligible'


def test_office_mentions_do_not_reject_canadian_vacancy():
    assert location_gate('We have offices in Spain and Poland.','Remote Canada')[0]=='eligible'
    assert location_gate('Lead teams remotely.','Remote')[0]=='unknown'


def test_javascript_page_does_not_erase_full_authoritative_feed_description():
    from job_search_assistant.searching.retrieval import enrich
    from job_search_assistant.ingestion import Page
    description='Lead backend engineering teams, develop engineers and deliver cloud infrastructure. '*3
    p=Posting('https://example.com/role/123','Example','Engineering Manager',
              description=description,location='Remote Canada',full_text=True,source='Greenhouse board:example')
    class Client:
        def get(self,url): return Page(url,'<html><title>Careers</title><p>Enable JavaScript to view this page.</p></html>')
    enrich(p,Client())
    assert p.description==description
    assert p.title=='Engineering Manager'
    assert p.full_text
    assert p.application_status=='unverified'


@pytest.mark.parametrize('title,location',[
    ('Senior Software Engineer','Remote Canada'),
    ('Senior Engineering Manager','Remote Spain'),
    ('Engineering Manager','Remote Poland'),
])
def test_export_cannot_leak_out_of_scope_jobs_into_verification(tmp_path,title,location):
    p=Posting('https://example.com/role/123','Example',title,location=location)
    xlsx,_=export(tmp_path,'scope',[(p,'verification','Possible duplicate',{})],[])
    wb=load_workbook(xlsx,data_only=True)
    assert wb['Needs Verification'].max_row==1
    assert wb['Recommended Jobs'].max_row==1
    assert wb['Rejected Jobs'].max_row==2
    wb.close()


@pytest.mark.parametrize('mode',['uncertain-ic','uncertain-location','retrieved-ic'])
def test_uncertain_duplicates_and_retrieved_titles_cannot_bypass_scope(tmp_path,monkeypatch,mode):
    from job_search_assistant.searching import engine
    tracker=tmp_path/'tracker.xlsx'
    wb=Workbook();ws=wb.active
    ws.append(['Company','Job title','Job posting (URL)']);wb.save(tracker);wb.close()
    p=Posting('https://example.com/role/123','Example',
              'Senior Software Engineer' if mode=='uncertain-ic' else 'Engineering Manager',
              location='Remote Poland' if mode=='uncertain-location' else 'Remote Canada')
    monkeypatch.setattr(engine,'capture_postings',lambda captures:[p])
    monkeypatch.setattr(engine,'load_service',lambda *a,**kw:SimpleNamespace(resume='Synthetic candidate'))
    monkeypatch.setattr(engine,'plan',lambda:[])
    class Reviewer:
        def review(self,p): pytest.fail('Out-of-scope job reached AI reviewer')
    monkeypatch.setattr(engine,'SearchReviewer',lambda *a:Reviewer())
    if mode.startswith('uncertain'):
        monkeypatch.setattr(engine.Store,'duplicate',lambda *a:('uncertain','Possible company/title duplicate'))
        monkeypatch.setattr(engine,'enrich',lambda *a:pytest.fail('Hard exclusion must precede uncertain duplicate'))
    else:
        def enrich(posting,client):
            posting.title='Senior Security Engineer'
            return posting
        monkeypatch.setattr(engine,'enrich',enrich)
    outputs=engine._run_search(tmp_path,workbook=tracker,progress=lambda s:None)
    wb=load_workbook(outputs[0],data_only=True)
    assert wb['Needs Verification'].max_row==1
    assert wb['Rejected Jobs'].max_row==2
    wb.close()
