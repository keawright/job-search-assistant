from dataclasses import replace
from datetime import date
from pathlib import Path
import json

from openpyxl import Workbook, load_workbook
import pytest

from job_search_assistant.assessment import AssessmentService, AssessmentError, CRITERIA
from job_search_assistant.ingestion import FetchError
from job_search_assistant.local_jds import LocalDescriptions
from job_search_assistant.models import Job
from job_search_assistant.workbook import update_workbook

TEXT = 'Example Engineering Manager\nLead distributed engineering teams and develop senior engineers. Own reliable cloud infrastructure and operational improvement programs.'
RESUME = 'Led two distributed cloud engineering teams. Delivered reliable infrastructure and developed senior engineers. Program leadership and operational improvements.'


def index(tmp_path, jobs, texts):
    folder = tmp_path / 'JDs'
    folder.mkdir(exist_ok=True)
    for name in texts:
        (folder / name).write_bytes(name.encode())
    return LocalDescriptions(tmp_path, jobs, reader=lambda p: (texts[p.name], 'Synthetic text'))


def test_identity_not_filename_and_ambiguous_matches_are_held(tmp_path):
    job = Job('one', 'Example', 'Engineering Manager', 'https://example.com/1')
    descriptions = index(tmp_path, [job], {'misnamed.pdf': TEXT})
    assert descriptions.find(job)[0].path.name == 'misnamed.pdf'
    other = replace(job, id='two', url='https://example.com/2')
    descriptions.jobs.append(other)
    assert descriptions.find(job)[0] is None
    assert 'multiple tracked jobs' in descriptions.find(job)[1][0]


def test_company_name_alone_is_not_a_match(tmp_path):
    job = Job('one', 'Example', 'Engineering Manager', 'https://example.com/1')
    descriptions = index(tmp_path, [job], {'Example.pdf': TEXT.replace('Engineering Manager', 'Sales Director')})
    assert descriptions.find(job)[0] is None


def test_changed_pdf_invalidates_extraction_cache(tmp_path):
    job = Job('one', 'Example', 'Engineering Manager', 'https://example.com/1')
    descriptions = index(tmp_path, [job], {'one.pdf': TEXT})
    first = descriptions.find(job)[0]
    (tmp_path / 'JDs/one.pdf').write_bytes(b'changed')
    descriptions = LocalDescriptions(tmp_path, [job], reader=lambda p: (TEXT + ' New requirements.', 'Synthetic'))
    second = descriptions.find(job)[0]
    assert first.sha256 != second.sha256
    assert 'New requirements.' in second.text


def test_unreadable_and_conflicting_pdfs_stay_unmatched(tmp_path):
    job = Job('one', 'Example', 'Engineering Manager', 'https://example.com/1')
    descriptions = index(tmp_path, [job], {'one.pdf': TEXT, 'two.pdf': TEXT + ' Different duties.'})
    assert descriptions.find(job)[0] is None
    descriptions = LocalDescriptions(tmp_path, [job], reader=lambda p: ('', 'Synthetic'))
    for p in (tmp_path / '.job-search-assistant-private/jd-text').glob('*.json'):
        p.write_text('{}')
    assert descriptions.find(job)[0] is None
    assert descriptions.errors


def test_reviewed_mapping_requires_exact_pdf_hash(tmp_path):
    job = Job('one', 'Example', 'Engineering Manager', 'https://example.com/1')
    descriptions = index(tmp_path, [job], {'one.pdf': TEXT.replace('Example', 'OCR damaged employer')})
    doc = descriptions.documents()[0]
    mapping = {'one.pdf': {'url': job.url, 'sha256': doc.sha256}}
    (tmp_path / 'JDs/matches.json').write_text(json.dumps(mapping))
    assert descriptions.find(job)[0] == doc
    mapping['one.pdf']['sha256'] = 'incorrect'
    (tmp_path / 'JDs/matches.json').write_text(json.dumps(mapping))
    assert descriptions.find(job)[0] is None
    assert 'changed' in descriptions.find(job)[1][0]


def test_invalid_mapping_does_not_silently_fall_back(tmp_path):
    job = Job('one', 'Example', 'Engineering Manager', 'https://example.com/1')
    descriptions = index(tmp_path, [job], {'one.pdf': TEXT})
    (tmp_path / 'JDs/matches.json').write_text('{broken')
    assert descriptions.find(job)[0] is None
    assert 'Invalid JDs/matches.json' in descriptions.find(job)[1][0]


@pytest.mark.parametrize('state', ['Assessment failed', 'Description needed', 'runtime failure', 'unknown result'])
def test_workbook_pdf_fallback_preserves_status_and_availability(tmp_path, monkeypatch, state):
    monkeypatch.setattr('job_search_assistant.local_jds.extract_pdf', lambda p: (TEXT, 'Synthetic'))
    # Patch constructor default explicitly (defaults bind at definition time).
    real_index = LocalDescriptions
    monkeypatch.setattr('job_search_assistant.workbook.LocalDescriptions',
                        lambda root, jobs, **kw: real_index(root, jobs, reader=lambda p: (TEXT, 'Synthetic'), **kw))
    (tmp_path / 'JDs').mkdir()
    (tmp_path / 'JDs/role.pdf').write_bytes(b'synthetic-pdf')
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Assistant queue'
    ws.append(['Company', 'Position title', 'Job posting', 'Application status', 'Assistant description',
               'Assistant assessment outcome', 'My notes'])
    ws.append(['Example', 'Engineering Manager', 'https://example.com/1', 'Not applied',
               'Old online description with unrelated company boilerplate. ' * 3, state, 'Keep'])
    ws.append(['Example', 'Engineering Director', 'https://example.com/2', 'Applied', 'Keep description', state, 'Protected'])
    wb.save(source); wb.close()
    class Provider:
        name, model = 'Synthetic', 'test'
        calls = 0
        def assess(self, resume, context, candidate):
            self.calls += 1
            if 'Lead distributed' not in context['description']:
                if state == 'unknown result':
                    return {**{key: {'rating': 'unknown', 'resume_quote': '', 'job_quote': '', 'rationale': 'Missing role'} for key in CRITERIA},
                            'strengths': [], 'gaps': [], 'regular_hands_on_coding': None, 'coding_quote': '',
                            'summary': 'No role evidence', 'uncertainties': []}
                raise AssessmentError('Assessment failed')
            return {**{key: {'rating': 'good', 'resume_quote': 'Led two distributed cloud engineering teams.',
                            'job_quote': 'Lead distributed engineering teams and develop senior engineers.',
                            'rationale': 'Relevant leadership'} for key in CRITERIA},
                    'strengths': [], 'gaps': [], 'regular_hands_on_coding': None, 'coding_quote': '',
                    'summary': 'Relevant leadership', 'uncertainties': []}
    def blocked(url):
        raise FetchError('Blocked')
    provider = Provider()
    service = AssessmentService(tmp_path / 'private', RESUME, provider)
    update_workbook(source, output, [], as_of=date(2026, 10, 28), assessment_service=service, fetcher=blocked)
    wb = load_workbook(output); it = wb['Assistant queue'].values; headers = next(it)
    rows = {r[headers.index('Application status')]: dict(zip(headers, r)) for r in it}; wb.close()
    row = rows['Not applied']
    assert row['Assistant assessment status'] == 'Complete'
    assert row['Assistant local description file'] == 'role.pdf'
    assert row['Assistant description'] == TEXT
    assert row['Assistant local description SHA256']
    assert row['My notes'] == 'Keep'
    assert row['Assistant local description method'] == 'Synthetic'
    assert row['Assistant location eligibility'] != 'eligible'
    assert 'availability' in row['Assistant review flags']
    assert rows['Applied']['Assistant description'] == 'Keep description'
    assert provider.calls == (2 if state in ('runtime failure', 'unknown result') else 1)
