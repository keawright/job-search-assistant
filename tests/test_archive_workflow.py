from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from job_search_assistant import cli
from job_search_assistant.archive import archive_older_workbooks, fingerprint
from job_search_assistant.workbook import update_workbook


def make_source(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Assistant queue"
    ws.append(["Company", "Position title", "Job posting", "Application status", "Applied Date", "Outcome Reason", "My notes"])
    ws.append(["Rejected first", "EM", "https://example.com/rejected", "Rejected", None, "My decision", "Keep rejection"])
    ws.append(["Needs review", "PM", "https://example.com/review", "Interviewing", None, None, "Custom status"])
    ws.append(["Applied missing URL", "Director", None, "Applied", "2026-10-14", None, "Recruiter"])
    ws.append(["Outstanding", "SEM", "https://example.com/outstanding", "Not applied", None, None, "Follow up"])
    ws.append(["Rejected missing URL", "PM", None, " rejected ", None, "Manual exclusion", "Keep exact status"])
    wb.create_sheet("Old history")["A1"] = "Preserve in archive"
    wb.save(path)
    wb.close()


def records(path):
    wb = load_workbook(path)
    ws = wb["Assistant queue"]
    values = list(ws.values)
    result = [dict(zip(values[0], row)) for row in values[1:]]
    names = wb.sheetnames
    wb.close()
    return names, result


def test_rejected_last_even_after_review_and_with_invalid_job(tmp_path):
    source, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    make_source(source)
    update_workbook(source, output, [], as_of=date(2026, 10, 15), offline=True)
    _, result = records(output)
    assert [r["Assistant section"] for r in result] == ["Applied", "Outstanding", "Review", "Rejected", "Rejected"]
    assert result[-1]["Application status"] == " rejected "
    assert result[-1]["Outcome Reason"] == "Manual exclusion"


def test_cli_archives_old_files_and_exports_one_current_tab(tmp_path, monkeypatch):
    source, intermediate = tmp_path / "latest.xlsx", tmp_path / "intermediate.xlsx"
    make_source(source)
    intermediate.write_bytes(b"Old version bytes")
    before = fingerprint(source)
    assert cli.main(["update", "--workbook", str(source), "--output", str(tmp_path / "next.xlsx"), "--offline"]) == 0
    assert sorted(p.name for p in tmp_path.glob("*.xlsx")) == ["next.xlsx"]
    assert fingerprint(tmp_path / "archive" / "latest.xlsx") == before
    assert (tmp_path / "archive" / "intermediate.xlsx").read_bytes() == b"Old version bytes"
    names, result = records(tmp_path / "next.xlsx")
    assert names == ["Assistant queue"]
    assert result[0]["Application status"] == "Applied"
    assert result[0]["Applied Date"] == "2026-10-14"
    # Manually update the current file and let the next command select it.
    wb = load_workbook(tmp_path / "next.xlsx")
    ws = wb.active
    headers = {c.value: c.column for c in ws[1]}
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, headers["Company"]).value == "Outstanding":
            ws.cell(row, headers["Application status"], "Rejected")
            ws.cell(row, headers["Outcome Reason"], "Manually ruled out today")
    wb.save(tmp_path / "next.xlsx")
    wb.close()
    monkeypatch.chdir(tmp_path)
    assert cli.main(["update", "--offline", "--url", "https://example.com/outstanding"]) == 0
    latest = list(tmp_path.glob("*.xlsx"))
    assert len(latest) == 1 and latest[0].name.startswith("jobs-updated-")
    _, result = records(latest[0])
    rejected = next(r for r in result if r["Company"] == "Outstanding")
    assert rejected["Application status"] == "Rejected"
    assert rejected["Outcome Reason"] == "Manually ruled out today"
    assert rejected["My notes"] == "Follow up"
    assert result[-1]["Assistant section"] == "Rejected"


def test_archive_collisions_never_replace_previous_backup(tmp_path):
    latest = tmp_path / "latest.xlsx"
    latest.write_bytes(b"current")
    (tmp_path / "older.xlsx").write_bytes(b"new backup")
    archive = tmp_path / "archive"
    archive.mkdir()
    (archive / "older.xlsx").write_bytes(b"existing backup")
    archived = archive_older_workbooks(latest)
    assert archived == [archive / "older-1.xlsx"]
    assert (archive / "older.xlsx").read_bytes() == b"existing backup"
    assert latest.read_bytes() == b"current"


def test_failed_update_does_not_archive_source(tmp_path):
    source = tmp_path / "latest.xlsx"
    make_source(source)
    before = fingerprint(source)
    assert cli.main(["update", "--workbook", str(source), "--output", str(tmp_path / "next.xlsx"), "--url", "file:///secret"]) == 2
    assert fingerprint(source) == before
    assert not (tmp_path / "archive").exists()


def test_concurrent_edits_prevent_archiving_newer_user_changes(tmp_path, monkeypatch, capsys):
    source = tmp_path / "latest.xlsx"
    make_source(source)
    original = cli.update_workbook
    def changed(*args, **kwargs):
        result = original(*args, **kwargs)
        wb = load_workbook(source)
        wb.active["G2"] = "A concurrent manual edit"
        wb.save(source)
        wb.close()
        return result
    monkeypatch.setattr(cli, "update_workbook", changed)
    assert cli.main(["update", "--workbook", str(source), "--output", str(tmp_path / "next.xlsx"), "--offline"]) == 2
    assert "Input workbook changed" in capsys.readouterr().err
    assert source.exists() and not (tmp_path / "archive").exists()
    wb = load_workbook(source)
    assert wb.active["G2"].value == "A concurrent manual edit"
    wb.close()


def test_default_source_refuses_ambiguous_workbooks(tmp_path, monkeypatch):
    make_source(tmp_path / "first.xlsx")
    make_source(tmp_path / "second.xlsx")
    monkeypatch.chdir(tmp_path)
    assert cli.main(["update", "--offline"]) == 2
    assert len(list(tmp_path.glob("*.xlsx"))) == 2


def test_cross_sheet_formula_stops_single_tab_export_without_archiving(tmp_path):
    source, output = tmp_path / "latest.xlsx", tmp_path / "next.xlsx"
    make_source(source)
    wb = load_workbook(source)
    wb.active["G2"] = "='Old history'!A1"
    wb.save(source)
    wb.close()
    assert cli.main(["update", "--workbook", str(source), "--output", str(output), "--offline"]) == 2
    assert source.exists() and not output.exists()
    assert cli.main(["update", "--workbook", str(source), "--output", str(output), "--offline", "--keep-history-tabs"]) == 0
    assert "Old history" in records(output)[0]
