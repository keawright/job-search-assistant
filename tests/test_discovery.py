from datetime import date
import json

from openpyxl import Workbook, load_workbook
import pytest

from job_search_assistant.discovery import Identity, find_alternative, linkedin_identity, search_postings
from job_search_assistant.ingestion import FetchError, Page
from job_search_assistant.workbook import update_workbook


LINK = "https://www.linkedin.com/jobs/view/123"
ALT = "https://example.com/jobs/123"
TITLE = "Example Co hiring Engineering Manager in Halifax, Canada | LinkedIn"


def posting(company="Example Co", title="Engineering Manager", location="Halifax Canada"):
    data = {"@type": "JobPosting", "title": title, "hiringOrganization": {"name": company}, "description": "Lead teams"}
    return '<script type="application/ld+json">' + json.dumps(data) + f'</script><main><p>{location}</p></main>'


@pytest.mark.parametrize("title,expected", [
    (TITLE, Identity("Example Co", "Engineering Manager", "Halifax, Canada")),
    ("Engineering Manager at Example Co — Canada | LinkedIn Jobs", Identity("Example Co", "Engineering Manager", "Canada")),
    ("Sign in | LinkedIn", None),
])
def test_identity(title, expected):
    assert linkedin_identity(title) == expected


def test_search_extracts_deduplicates_and_limits_results():
    def fetch(url):
        assert "Example%20Co" in url
        return Page(url, ''.join(f'<a class="result__a" href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fexample.com%2F{i}">job</a>' for i in [1, 1, 2, 3, 4]))
    assert search_postings(Identity("Example Co", "EM"), fetch) == [f"https://example.com/{i}" for i in [1, 2, 3]]


@pytest.mark.parametrize("html", [posting(company="Another Co"), posting(title="Director"), posting(location="Toronto Canada"), '<main>This position is closed</main>'])
def test_wrong_or_unverifiable_result_is_rejected(html):
    identity = linkedin_identity(TITLE)
    assert find_alternative(identity, lambda url: Page(url, html), lambda *args: [ALT]) is None


def test_ambiguous_results_are_rejected():
    assert find_alternative(linkedin_identity(TITLE), lambda url: Page(url, posting()), lambda *args: [ALT, ALT + "other"]) is None


@pytest.mark.parametrize("existing", [True, False])
def test_update_repairs_placeholders_retains_original_and_reuses_alternate(tmp_path, existing):
    source, output, again = [tmp_path / name for name in ["source.xlsx", "output.xlsx", "again.xlsx"]]
    wb = Workbook()
    ws = wb.active
    ws.title = "Assistant queue"
    ws.append(["Company", "Position title", "Job posting", "Application status", "My notes"])
    if existing:
        ws.append(["Unknown company", TITLE, LINK, "Not applied", "Keep"])
    wb.save(source)
    wb.close()
    calls = []
    def fetch(url):
        calls.append(url)
        if url == LINK:
            return Page(url, f'<title>{TITLE}</title><main>Job page</main>')
        if url.startswith("https://html.duckduckgo.com/"):
            return Page(url, f'<a class="result__a" href="{ALT}">Job</a>')
        assert url == ALT
        return Page(url, posting())
    update_workbook(source, output, [] if existing else [LINK], as_of=date(2026, 10, 20), fetcher=fetch)
    wb = load_workbook(output)
    data = list(wb["Assistant queue"].values)
    row = dict(zip(data[0], data[1]))
    wb.close()
    assert row["Company"] == "Example Co"
    assert row["Position title"] == "Engineering Manager"
    assert row["Job posting"] == LINK
    assert row["Assistant alternate URL"] == ALT
    assert row["Apply order"] == "Needs assessment"
    if existing:
        assert row["My notes"] == "Keep"
    assert "Source authority" in row["Assistant review flags"]
    calls.clear()
    update_workbook(output, again, [], as_of=date(2026, 10, 20), fetcher=fetch)
    assert calls == [ALT]


def test_blocked_search_does_not_lose_extracted_identity(tmp_path):
    from job_search_assistant.workbook import Record, _fetch
    record = Record({"Company": "Unknown company", "Position title": TITLE, "Application status": "Not applied"}, row=2)
    def fetch(url):
        if url == LINK:
            raise FetchError("HTTP 403")
        raise FetchError("Search blocked")
    assert not _fetch(record, LINK, tmp_path, date(2026, 10, 20), fetch)
    assert record.values["Company"] == "Example Co"
    assert "Search blocked" in record.values["Assistant discovery notes"]
    assert record.values["Application status"] == "Not applied"
    # The saved identity allows another search even after the fallback page title
    # has been replaced with a clean role and LinkedIn remains inaccessible.
    assert not _fetch(record, LINK, tmp_path, date(2026, 10, 21), fetch)
    assert "Search blocked" in record.values["Assistant discovery notes"]


def test_discovery_preserves_manual_company_and_title(tmp_path):
    from job_search_assistant.workbook import Record, _fetch
    record = Record({"Company": "My company label", "Position title": "My role label", "Application status": "Not applied"}, row=2)
    def fetch(url):
        return Page(url, f'<title>{TITLE}</title>') if url == LINK else Page(url, posting())
    assert _fetch(record, LINK, tmp_path, date(2026, 10, 20), fetch, lambda *args: [ALT])
    assert record.values["Company"] == "My company label"
    assert record.values["Position title"] == "My role label"
