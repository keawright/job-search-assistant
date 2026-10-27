from dataclasses import asdict
from pathlib import Path
import json
from openpyxl import Workbook
import pytest

from job_search_assistant.searching.normalization import Posting, canonical, identities
from job_search_assistant.searching.storage import Store
from job_search_assistant.searching.eligibility import location_gate
from job_search_assistant.searching.evaluation import decide
from job_search_assistant.searching.retrieval import PublicClient
from job_search_assistant.ingestion import Page, FetchError
from job_search_assistant.assessment import CRITERIA, AssessmentError


def test_canonical_ids_link_custom_employer_and_ats_urls():
    assert identities('https://example.com/job?gh_jid=123&utm_source=feed') & identities('https://job-boards.greenhouse.io/demo/jobs/123?gh_src=tracking') == {'greenhouse:123'}
    assert canonical('https://example.com/job/?gh_src=test&utm_campaign=foo')=='https://example.com/job'
    assert canonical('https://example.com/job?id=3') != canonical('https://example.com/job?id=4')


def test_exclusions_include_hidden_archived_sheets_and_hyperlinks(tmp_path):
    p=tmp_path/'input.xlsx'; wb=Workbook(); ws=wb.active
    ws.append(['Company','Position title','Job posting (URL)','Application status'])
    ws.append(['Example','Engineering Manager','https://example.com/job?gh_jid=123','Closed'])
    ws=wb.create_sheet('History'); ws.sheet_state='hidden'
    ws.append(['Company','Job title','Other source links (URL)'])
    ws.append(['Other','Technical Program Manager','Apply here']); ws['C2'].hyperlink='https://example.com/other'
    wb.save(p); wb.close()
    store=Store(tmp_path/'db.sqlite3'); assert store.import_workbook(p)==2
    assert store.duplicate(Posting('https://job-boards.greenhouse.io/demo/jobs/123'))[0]=='duplicate'
    assert store.duplicate(Posting('https://example.com/other'))[0]=='duplicate'
    assert store.duplicate(Posting('https://example.com/new','Example','Engineering Manager - Canada'))[0]=='duplicate'
    store.close()


@pytest.mark.parametrize('text,location,expected',[
    ('Work remotely anywhere in Canada.','Remote Canada','eligible'),
    ('Remote US-only role','','ineligible'),
    ('We cannot hire employees in Nova Scotia.','','ineligible'),
    ('Remote within Ontario only.','','ineligible'),
    ('Remote work supported','Toronto, Remote','unknown'),
    ('Hybrid in Halifax, two days per week','','eligible'),
    ('Must attend the office in Toronto three days weekly.','','ineligible'),
    ('Canada-based team.','','unknown'),
])
def test_location_filter(text,location,expected):
    assert location_gate(text,location)[0]==expected


def review_fixture():
    resume='Led distributed engineering teams delivering reliable cloud platforms and cross-team programs.'
    text='Lead engineering teams and deliver reliable cloud platforms. Remote anywhere in Canada. Apply for this job. Leadership and program delivery experience required.'
    p=Posting('https://example.com/job','Example','Engineering Manager',text,'Remote Canada',source='Employer public page',full_text=True,application_evidence='Apply for this job')
    from datetime import date
    p.application_status='accepting';p.application_checked_on=date.today().isoformat()
    p.application_verification_url=p.url;p.application_verification_method='live_employer_application_form'
    p.application_verification_evidence='Current employer form exposes resume, email and applicant name fields'
    p.application_verification_fresh=True
    a={**{k:dict(rating='good',resume_quote='Led distributed engineering teams',job_quote='Lead engineering teams',rationale='Demonstrated comparable leadership scope.') for k in CRITERIA},
       'strengths':['Platform leadership'],'gaps':[],'regular_hands_on_coding':None,'coding_quote':'',
       'summary':'Relevant platform delivery and people leadership.','uncertainties':['Confirm working hours.']}
    gate=lambda status,quote:dict(status=status,quote=quote,reason='Supported by supplied evidence.')
    r=dict(assessment=a,location=gate('eligible','Remote anywhere in Canada'),vacancy=gate('open','Apply for this job'),
           requirements=gate('met','Leadership and program delivery experience required'),primary_coding=gate('no','Lead engineering teams and deliver reliable cloud platforms'),work_life_balance='Unknown',compensation='Not published')
    return p,r,resume


def test_recommendation_requires_every_material_gate():
    p,r,resume=review_fixture()
    assert decide(p,r,resume)==('recommended','Good Fit')
    r['requirements']['status']='unknown'; r['requirements']['quote']=''
    assert decide(p,r,resume)[0]=='verification'
    r['requirements']['status']='gap'; r['requirements']['quote']='Leadership and program delivery experience required'
    assert decide(p,r,resume)[0]=='rejected'


def test_search_snippets_and_unsupported_quotes_never_recommended():
    p,r,resume=review_fixture(); p.full_text=False
    assert decide(p,r,resume)[0]=='verification'
    p.full_text=True; r['location']['quote']='Invented nationwide permission'
    with pytest.raises(AssessmentError): decide(p,r,resume)


def test_ai_open_verdict_cannot_replace_independent_current_application_check():
    p,r,resume=review_fixture()
    p.application_status='unverified';p.application_verification_fresh=False
    assert r['vacancy']['status']=='open'
    state,reason=decide(p,r,resume)
    assert state=='verification'
    assert 'independently verified today' in reason


def test_persistent_history_preserves_decisions_and_excludes_reported(tmp_path):
    p,r,resume=review_fixture(); store=Store(tmp_path/'db.sqlite3')
    store.record(p,'recommended','Good Fit',r,'run1')
    store.db.execute('UPDATE findings SET reported_on=?',('run1',)); store.db.commit(); store.close()
    store=Store(tmp_path/'db.sqlite3')
    assert store.duplicate(p)[0]=='duplicate'
    alias=Posting('https://other.example/job','Example','Engineering Manager')
    assert store.duplicate(alias)[0]=='duplicate'
    assert store.db.execute('SELECT count(*) FROM events').fetchone()[0]==1
    store.close()


def test_robots_denial_and_failed_host_do_not_fetch_jobs():
    calls=[]
    def fetch(url,**kwargs):
        calls.append(url); return Page(url,'User-agent: *\nDisallow: /jobs')
    client=PublicClient(fetcher=fetch,delay=0)
    with pytest.raises(FetchError,match='disallows'): client.get('https://example.com/jobs/1')
    assert calls==['https://example.com/robots.txt']
    def blocked(url,**kwargs): raise FetchError('HTTP 403')
    client=PublicClient(fetcher=blocked,delay=0)
    for _ in range(2):
        with pytest.raises(FetchError): client.get('https://example.com/jobs/1')
    assert client.calls==1


@pytest.mark.parametrize("first,second", [
    ("https://example.zohorecruit.com/jobs/Careers/12345/", "https://example.zohorecruit.com/recruit/PortalDetail.na?jobid=12345&source=feed"),
    ("https://example.wd1.myworkdayjobs.com/en-US/External/job/Manager_R-12345", "https://example.wd1.myworkdayjobs.com/External/job/Manager_R-12345-1"),
])
def test_other_ats_ids_survive_alternate_urls(first,second):
    assert identities(first) & identities(second)


def test_grounded_secondary_role_exports_once_and_never_edits_tracker(tmp_path):
    from job_search_assistant.searching.engine import run_search
    from job_search_assistant.searching.evaluation import review_hash
    p,r,resume=review_fixture(); p.title="Senior Project Manager (Contractual)"
    resume += " Mentored senior engineers and managed delivery risks."
    private=tmp_path/".job-search-assistant-private"; private.mkdir()
    (private/"resume.txt").write_text(resume,encoding="utf-8")
    (private/"assessment-config.json").write_text(json.dumps({"enabled":True,"resume_path":".job-search-assistant-private/resume.txt"}))
    source=tmp_path/"jobs.xlsx"; wb=Workbook();ws=wb.active
    ws.append(["Company","Position title","Job posting (URL)"])
    ws.append(["Existing","Another role","https://existing.example/job"])
    wb.save(source);wb.close(); before=source.read_bytes()
    bundle=private/"reviews.json"
    bundle.write_text(json.dumps([dict(posting=asdict(p),review=r,input_hash=review_hash(p,resume))]),encoding="utf-8")
    first=run_search(tmp_path,reviews=bundle,offline=True,progress=lambda *args:None)
    from openpyxl import load_workbook
    report=load_workbook(first[0],data_only=True)
    assert report["Recommended Jobs"].max_row==2
    assert report["Needs Verification"].max_row==1
    report.close()
    second=run_search(tmp_path,reviews=bundle,offline=True,progress=lambda *args:None)
    report=load_workbook(second[0],data_only=True)
    assert report["Recommended Jobs"].max_row==1
    report.close()
    assert source.read_bytes()==before
    assert len(list(tmp_path.glob("*.xlsx")))==1



def test_report_sanitizes_search_controls_and_formula_injection(tmp_path):
    from job_search_assistant.searching.reporting import export
    p,r,resume=review_fixture();p.title="=unsafe"+chr(0)+"title"
    xlsx,md=export(tmp_path,"synthetic",[(p,"verification","Need employer confirmation",r)],[])
    from openpyxl import load_workbook
    wb=load_workbook(xlsx);cell=wb["Needs Verification"].cell(2,2)
    assert cell.value=="'=unsafetitle"
    assert cell.data_type=='s'
    assert wb["Needs Verification"].cell(2,12).hyperlink.target==p.url
    wb.close()


def test_redirect_checks_destination_robots_before_following():
    calls=[]
    def fetch(url,**kwargs):
        calls.append(url)
        if url=='https://one.example/robots.txt': return Page(url,'User-agent: *\nAllow: /')
        if url=='https://two.example/robots.txt': return Page(url,'User-agent: *\nDisallow: /')
        kwargs['redirect_check']('https://two.example/jobs/1')
        pytest.fail('Redirect followed despite destination robots denial')
    client=PublicClient(fetcher=fetch,delay=0)
    with pytest.raises(FetchError,match='disallows'): client.get('https://one.example/jobs/1')
    assert calls==['https://one.example/robots.txt','https://one.example/jobs/1','https://two.example/robots.txt']


@pytest.mark.parametrize('html',[
    '<main><button>Sorry, this job has expired</button></main>',
    '<main><h2>The job you are trying to apply for is no longer available.</h2></main>',
])
def test_expired_employer_application_controls_are_closure_evidence(html):
    from job_search_assistant.ingestion import closure_evidence
    assert closure_evidence(html)


@pytest.mark.parametrize('rating,expected',[('strong','recommended'),('good','recommended'),('stretch','rejected'),('unsuitable','rejected'),('unknown','verification')])
def test_fit_classifications_never_promote_stretch_or_unknown(rating,expected):
    p,r,resume=review_fixture()
    for key in CRITERIA: r['assessment'][key]['rating']=rating
    assert decide(p,r,resume)[0]==expected
    if rating=='strong': assert decide(p,r,resume)[1]=='Strong Fit'


def test_primary_coding_rejected_even_with_strong_leadership():
    p,r,resume=review_fixture()
    p.description+=' Write production code daily as your primary responsibility.'
    r['primary_coding'].update(status='yes',quote='Write production code daily as your primary responsibility')
    assert decide(p,r,resume)[0]=='rejected'


def test_long_provincial_restriction_overrides_remote_canada_label():
    text="Candidates for this remote work opportunity must be based in either Alberta, British Columbia or Ontario. If you are based in another location within Canada, we aren't able to hire in your location at this time."
    assert location_gate(text,'Remote Canada')[0]=='ineligible'


@pytest.mark.parametrize('title',['Manager, Cloud Infrastructure','Director, DevOps','Head of Backend','Senior Project Manager (Contractual)'])
def test_transferable_title_variations_are_routed_for_full_review(title):
    from job_search_assistant.searching.eligibility import title_gate
    assert title_gate(title)


@pytest.mark.parametrize('platform',['greenhouse','lever','ashby'])
def test_public_feed_adapters_preserve_role_text_without_inventing_post_date(platform):
    from job_search_assistant.searching.discovery import ats_board
    description='Lead distributed platform teams and own delivery of reliable cloud services. '*3
    class Client:
        def get(self,url,json_response=False):
            assert json_response
            if platform=='greenhouse':
                return {'jobs':[dict(absolute_url='https://example.com/role',title='Engineering Manager',content='<p>'+description+'</p>',location={'name':'Remote Canada'},updated_at='2026-10-27')]}
            if platform=='lever':
                return [dict(hostedUrl='https://example.com/role',text='Engineering Manager',descriptionPlain=description,categories={'location':'Remote Canada'},lists=[])]
            return {'jobs':[dict(jobUrl='https://example.com/role',title='Engineering Manager',descriptionPlain=description,location='Canada',isRemote=True,publishedAt='2026-10-27')]}
    p=list(ats_board(platform,'example',Client()))[0]
    assert 'reliable cloud services' in p.description
    assert p.full_text
    assert p.posted==''
    assert p.application_evidence=='Listed in public employer ATS feed'
