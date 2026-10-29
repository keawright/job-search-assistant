from dataclasses import replace
from datetime import date,timedelta

import pytest
from openpyxl import load_workbook

from job_search_assistant.searching.normalization import Posting
from job_search_assistant.searching.verification import inspect_connected,current_application_verified,inspect_live
from job_search_assistant.searching.reporting import export
from job_search_assistant.ingestion import Page


@pytest.mark.parametrize('notice', ['This job is closed.', 'This position is no longer active.',
                                  'This job posting has been removed.', 'This role has already been filled.'])
def test_closed_notice_detected_without_job_description(notice):
    p=Posting('https://example.com/jobs/123','Example','Engineering Manager')
    inspect_connected(p,connected(p.url,body=notice))
    assert p.application_status=='closed'


@pytest.mark.parametrize('error,closed', [('HTTP 404', True), ('HTTP 410', True),
                                        ('HTTP 403', False), ('HTTP 429', False),
                                        ('Robots/access check unavailable: HTTP 404', False)])
def test_removed_vacancy_distinguished_from_access_failure(error,closed):
    from job_search_assistant.searching.verification import verify_live
    from job_search_assistant.ingestion import FetchError
    p=Posting('https://example.com/jobs/123','Example','Engineering Manager')
    class Client:
        def get(self,url): raise FetchError(error)
    if closed:
        verify_live(p,Client())
        assert p.application_status=='closed'
    else:
        with pytest.raises(FetchError): verify_live(p,Client())
        assert p.application_status=='unverified'


def test_closed_posting_exported_as_rejected_without_fit_review(tmp_path):
    p=Posting('https://example.com/jobs/123','Example','Engineering Manager')
    inspect_live(p,Page(p.url,'<main><h2>This job is closed.</h2></main>'))
    xlsx,_=export(tmp_path,'closed',[(p,'verification','Description missing',{})],[])
    wb=load_workbook(xlsx,data_only=True)
    assert wb['Needs Verification'].max_row==1
    assert wb['Recommended Jobs'].max_row==1
    assert wb['Rejected Jobs'].max_row==2
    wb.close()


def test_third_party_404_and_generic_career_404_do_not_establish_job_closure():
    from job_search_assistant.searching.verification import mark_removed
    p=Posting('https://example.com/jobs/123','Example','Engineering Manager')
    assert not mark_removed(p,'https://linkedin.com/jobs/view/123','HTTP 404')
    assert not mark_removed(p,'https://example.com/careers','HTTP 404')

FORM='Apply for this job\nFirst Name[Input]\nLast Name[Input]\nEmail[Input]\nResume/CV[Input]\nSubmit application'


def connected(url,age='today',body=FORM):
    return f'Job ({url})\nCrawled: {age}; Content type: text/html\nL0: '+body.replace('\n','\nL1: ')


def test_current_employer_form_is_required_not_an_apply_button():
    p=Posting('https://job-boards.greenhouse.io/example/jobs/123','Example','Manager')
    inspect_connected(p,connected(p.url,body='Engineering Manager. Apply now'))
    assert not current_application_verified(p)
    inspect_connected(p,connected(p.url))
    assert current_application_verified(p)


@pytest.mark.parametrize('age',['yesterday','last week','3 months ago'])
def test_old_snapshot_cannot_prove_current_acceptance(age):
    p=Posting('https://job-boards.greenhouse.io/example/jobs/123','Example','Manager')
    inspect_connected(p,connected(p.url,age))
    assert p.application_status=='unverified'
    assert not current_application_verified(p)


def test_proof_must_be_for_same_job_and_checked_today():
    p=Posting('https://job-boards.greenhouse.io/example/jobs/123','Example','Manager')
    inspect_connected(p,connected(p.url))
    p.application_verification_url='https://job-boards.greenhouse.io/example/jobs/456'
    assert not current_application_verified(p)
    p.application_verification_url=p.url
    p.application_checked_on=(date.today()-timedelta(days=1)).isoformat()
    assert not current_application_verified(p)


def test_third_party_forms_do_not_establish_employer_status():
    p=Posting('https://wellfound.com/jobs/123','Example','Manager')
    inspect_connected(p,connected(p.url))
    assert p.application_status=='unverified'
    assert not current_application_verified(p)


def test_verified_closure_overrides_any_form_fields():
    p=Posting('https://job-boards.greenhouse.io/example/jobs/123','Example','Manager')
    inspect_connected(p,connected(p.url,body='This job is closed.\n'+FORM))
    assert p.application_status=='closed'
    assert not current_application_verified(p)


def test_conflicting_current_employer_pages_are_held_for_review():
    from job_search_assistant.searching.verification import choose_verification
    p=Posting('https://job-boards.greenhouse.io/example/jobs/123','Example','Manager')
    inspect_connected(p,connected(p.url))
    closed=replace(p)
    inspect_connected(closed,connected(p.url,body='This job is closed.'))
    result=choose_verification([p,closed])
    assert result.application_status=='unverified'
    assert not current_application_verified(result)


def test_export_independently_withholds_unverified_recommendations(tmp_path):
    p=Posting('https://example.com/job','Example','Engineering Manager',application_evidence='Apply for this job')
    review={'assessment':{'summary':'Relevant experience','uncertainties':[]}}
    xlsx,_=export(tmp_path,'gate',[(p,'recommended','Good Fit',review)],[])
    wb=load_workbook(xlsx,data_only=True)
    assert wb['Recommended Jobs'].max_row==1
    assert wb['Needs Verification'].max_row==2
    assert 'not verified' in wb['Needs Verification'].cell(2,14).value
    wb.close()


def test_live_form_is_dated_and_source_bound():
    p=Posting('https://example.com/job','Example','Manager')
    inspect_live(p,Page(p.url,'<main><h1>Manager</h1><form><h2>Apply for this position</h2><label>Full Name</label><label>Email</label><label>Resume/CV</label></form></main>'))
    assert current_application_verified(p)


def test_verification_metadata_does_not_invalidate_unchanged_fit_evidence():
    p=Posting('https://example.com/job','Example','Manager',description='Lead engineering teams')
    old=p.content_hash
    inspect_connected(p,connected(p.url))
    assert p.content_hash==old


def test_live_verification_follows_same_job_application_link_only():
    from job_search_assistant.searching.verification import verify_live
    url='https://jobs.lever.co/example/11111111-2222-3333-4444-555555555555'
    p=Posting(url,'Example','Manager');calls=[]
    class Client:
        def get(self,target):
            calls.append(target)
            if target==url:return Page(target,'<main><h1>Manager</h1><a href="'+url+'/apply">Apply for this job</a></main>')
            return Page(target,'<main><h1>Manager</h1><form><h2>Submit your application</h2><label>Full Name</label><label>Email</label><label>Resume/CV</label></form></main>')
    verify_live(p,Client())
    assert calls==[url,url+'/apply']
    assert current_application_verified(p)
