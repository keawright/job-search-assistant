from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
import pytest

from job_search_assistant.workbook import Record, _age_applied, update_workbook
from openpyxl.utils.datetime import WINDOWS_EPOCH


@pytest.mark.parametrize("applied,as_of,expected", [
    (date(2026, 9, 20), date(2026, 10, 19), "Applied"),
    (date(2026, 9, 20), date(2026, 10, 20), "Elapsed-Rejected"),
    (date(2026, 8, 20), date(2026, 10, 20), "Elapsed-Rejected"),
    (date(2026, 1, 31), date(2026, 2, 28), "Elapsed-Rejected"),
    (date(2028, 1, 31), date(2028, 2, 28), "Applied"),
    (date(2028, 1, 31), date(2028, 2, 29), "Elapsed-Rejected"),
    (date(2025, 12, 31), date(2026, 1, 31), "Elapsed-Rejected"),
    ("Oct-9-2026", date(2026, 11, 9), "Elapsed-Rejected"),
    ("2026-09-20", date(2026, 10, 20), "Elapsed-Rejected"),
    (None, date(2026, 10, 20), "Applied"),
    ("09/10/2026", date(2026, 10, 20), "Applied"),
    (date(2027, 1, 1), date(2026, 10, 20), "Applied"),
])
def test_one_calendar_month(applied, as_of, expected):
    record = Record({"Application status": "Applied", "Applied Date": applied})
    _age_applied(record, WINDOWS_EPOCH, as_of)
    assert record.values["Application status"] == expected
    assert record.values["Applied Date"] == applied
    if applied is None or applied == "09/10/2026" or applied == date(2027, 1, 1):
        assert "aging skipped" in record.flags[0]


def test_aging_persists_moves_to_bottom_and_preserves_user_fields(tmp_path):
    source, output, again = [tmp_path / name for name in ["source.xlsx", "output.xlsx", "again.xlsx"]]
    wb = Workbook()
    ws = wb.active
    ws.title = "Assistant queue"
    ws.append(["Company", "Position title", "Job posting", "Application status", "Applied Date", "My notes", "Outcome Reason"])
    for name, status, applied in [("Old", "Applied", date(2026, 9, 20)), ("Recent", "Applied", date(2026, 10, 1)), ("Rejected", "Rejected", None), ("Interview", "Interviewing", date(2026, 1, 1))]:
        ws.append([name, "EM", None, status, applied, "Keep notes", "Keep outcome"])
    ws["D2"].comment = Comment("Keep annotation", "User")
    wb.save(source)
    wb.close()
    def forbidden(url):
        pytest.fail("Protected status must not be fetched")
    for input_path, target in [(source, output), (output, again)]:
        update_workbook(input_path, target, [], as_of=date(2026, 10, 20), fetcher=forbidden)
        wb = load_workbook(target)
        ws = wb["Assistant queue"]
        data = list(ws.values)
        rows = [dict(zip(data[0], values)) for values in data[1:]]
        assert [row["Company"] for row in rows] == ["Interview", "Recent", "Rejected", "Old"]
        old = rows[-1]
        assert old["Application status"] == old["Assistant section"] == "Elapsed-Rejected"
        assert old["Apply order"] == "-"
        assert old["Applied Date"].date() == date(2026, 9, 20)
        assert old["My notes"] == "Keep notes"
        assert old["Outcome Reason"] == "Keep outcome"
        assert "not a confirmed" in old["Assistant priority reasons"]
        assert ws.cell(ws.max_row, data[0].index("Application status") + 1).comment.text == "Keep annotation"
        wb.close()
