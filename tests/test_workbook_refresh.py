from datetime import date
import json

from openpyxl import Workbook, load_workbook
import pytest

from job_search_assistant.ingestion import FetchError, Page, closure_evidence
from job_search_assistant.workbook import update_workbook

TODAY = date(2026, 10, 15)


def test_normal_update_closes_existing_and_new_and_does_not_recheck(tmp_path):
    original, output, again = (tmp_path / name for name in ("input.xlsx", "output.xlsx", "again.xlsx"))
    source(original)
    calls = []
    def closed(url):
        calls.append(url)
        return Page(url, '<main><h1>This position has been filled.</h1></main>')
    update_workbook(original, output, ["https://example.com/new"], as_of=TODAY, fetcher=closed)
    assert len(calls) == 2
    values = get_values(output)
    assert values["Application status"] == "Closed"
    assert values["Assistant section"] == "Closed"
    assert "filled" in values["Assistant closure evidence"]
    assert values["Comment"] == "Do not edit"
    def forbidden(url):
        pytest.fail("Closed job fetched again")
    update_workbook(output, again, [], as_of=TODAY, refresh=True, fetcher=forbidden)
    assert get_values(again)["Application status"] == "Closed"


@pytest.mark.parametrize("status", ["Applied", "Rejected", "Interviewing", "Closed", "Outstanding"])
def test_other_statuses_never_fetched_even_with_refresh(tmp_path, status):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    source(original)
    wb = load_workbook(original)
    wb.active["D2"] = status
    wb.save(original)
    wb.close()
    def forbidden(url):
        pytest.fail("Protected status fetched")
    update_workbook(original, output, [], as_of=TODAY, refresh=True, fetcher=forbidden)
    assert get_values(output)["Application status"] == status


@pytest.mark.parametrize("status", ["Not applied", None])
def test_inaccessible_existing_job_keeps_status_and_records_issue(tmp_path, status):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    source(original)
    wb = load_workbook(original)
    wb.active["D2"] = status
    wb.save(original)
    wb.close()
    def blocked(url):
        raise FetchError("HTTP 403")
    update_workbook(original, output, [], as_of=TODAY, fetcher=blocked)
    values = get_values(output)
    assert values["Application status"] == status
    assert "403" in values["Assistant fetch issue"]
    assert "403" in values["Assistant review flags"]


@pytest.mark.parametrize("html", [
    '<main><h1>Sign in</h1></main>',
    '<script>This job is closed</script><main><p>Engineer</p></main>',
    '<footer><p>This job is closed</p></footer><main>Engineer</main>',
    '<main><p>Applications close tomorrow</p></main>',
    '<script type="application/ld+json">{"@type":"JobPosting","validThrough":"2020-01-01"}</script>',
])
def test_uncertain_pages_are_not_closure_evidence(html):
    assert closure_evidence(html) is None


def test_apply_order_is_regenerated_sorted_and_stable_on_rerun(tmp_path):
    original, output, again = (tmp_path / name for name in ("input.xlsx", "output.xlsx", "again.xlsx"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Assistant queue"
    ws.append(["Company", "Position title", "Job posting", "Application status", "Apply order", "Manual priority", "Leadership fit", "Assessment evidence"])
    for name, status, old_order, manual, url in [
        ("Rejected", "Rejected", 1, None, None),
        ("Later", "Not Applied", 2, None, "https://example.com/later"),
        ("Review", "Not Applied", 3, None, None),
        ("First", "Not Applied", 99, 1, "https://example.com/first"),
        ("Closed", "Closed", 4, None, "https://example.com/closed"),
        ("Applied", "Applied", 5, None, None),
    ]:
        ws.append([name, "EM", url, status, old_order, manual, "good", "Comparable team leadership and delivery experience against the role's team-management requirement."])
    wb.save(original)
    wb.close()
    for source_path, target in [(original, output), (output, again)]:
        update_workbook(source_path, target, [], as_of=TODAY, offline=True)
        wb = load_workbook(target)
        data = list(wb["Assistant queue"].values)
        records = [dict(zip(data[0], row)) for row in data[1:]]
        wb.close()
        assert [r["Company"] for r in records] == ["Applied", "First", "Later", "Review", "Closed", "Rejected"]
        assert [r["Apply order"] for r in records] == ["-", 1, 2, "Needs assessment", "-", "-"]
        assert records[3]["Assistant section"] == "Review"
        assert "Missing posting URL" in records[3]["Assistant review flags"]
        assert [r["Application status"] for r in records] == ["Applied", "Not Applied", "Not Applied", "Not Applied", "Closed", "Rejected"]


def test_interviewing_is_first_and_preserved_on_rerun(tmp_path):
    original, output, again = (tmp_path / name for name in ("input.xlsx", "output.xlsx", "again.xlsx"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Assistant queue"
    ws.append(["Company", "Position title", "Job posting", "Application status", "Applied Date", "Comment"])
    for name, status, url in [
        ("Applied", "Applied", None),
        ("Interview one", "Interviewing", None),
        ("Outstanding", "Not Applied", "https://example.com/job"),
        ("Interview two", " interviewing ", "https://example.com/interview"),
        ("Rejected", "Rejected", None),
    ]:
        ws.append([name, "EM", url, status, "2020-01-01" if "interview" in status.lower() else None, "Keep my notes"])
    wb.save(original)
    wb.close()
    for source_path, target in [(original, output), (output, again)]:
        update_workbook(source_path, target, [], as_of=TODAY, offline=True)
        wb = load_workbook(target)
        data = list(wb["Assistant queue"].values)
        records = [dict(zip(data[0], row)) for row in data[1:]]
        wb.close()
        assert [r["Company"] for r in records[:3]] == ["Interview one", "Interview two", "Applied"]
        assert [r["Application status"] for r in records[:2]] == ["Interviewing", " interviewing "]
        for record in records[:2]:
            assert record["Assistant section"] == "Interviewing"
            assert record["Apply order"] == "-"
            assert record["Applied Date"] == "2020-01-01"
            assert record["Comment"] == "Keep my notes"


def source(path, extra_headers=(), extra_values=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "Application priorities"
    ws.append(["Company", "Position title", "Job posting", "Application status", "Comment", *extra_headers])
    ws.append(["Original Co", "Original Title", "https://example.com/1", "Not applied", "Do not edit", *extra_values])
    wb.save(path)
    wb.close()


def get_values(path):
    wb = load_workbook(path)
    ws = wb["Assistant queue"]
    data = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
    wb.close()
    return data


def test_refresh_never_overwrites_manual_fields_and_retains_full_description(tmp_path):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    source(original, ("Deadline",), (date(2026, 10, 30),))
    description = "Original job description. " * 1600
    def fetch(url):
        data = {"@type": "JobPosting", "title": "Fetched Title", "hiringOrganization": {"name": "Fetched Co"}, "description": description, "validThrough": "2026-10-20"}
        return Page(url, '<script type="application/ld+json">' + json.dumps(data) + '</script>')
    update_workbook(original, output, [], as_of=TODAY, refresh=True, fetcher=fetch)
    values = get_values(output)
    assert values["Company"] == "Original Co" and values["Position title"] == "Original Title"
    assert values["Comment"] == "Do not edit"
    assert values["Deadline"].date() == date(2026, 10, 30)
    assert len(values["Assistant description"]) == 30000
    snapshot = json.loads(__import__('pathlib').Path(values["Assistant snapshot"]).read_text(encoding="utf-8"))
    assert snapshot["extracted"]["description"] == description.strip()


@pytest.mark.parametrize("headers,values", [(("Vacancy status",), ("Closed on LinkedIn",)),
    (("Vacancy status",), ("Apparently closed — not on current board",)),
        ])
def test_legacy_exclusions_do_not_silently_become_actionable(tmp_path, headers, values):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    source(original, headers, values)
    update_workbook(original, output, [], as_of=TODAY, offline=True)
    assert get_values(output)["Assistant section"] == "Review"


def test_passed_date_excluded_and_original_formula_preserved(tmp_path):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    source(original, ("Deadline", "Calculation"), (date(2026, 10, 14), "=LEN(A2)"))
    update_workbook(original, output, [], as_of=TODAY, offline=True)
    values = get_values(output)
    assert values["Assistant section"] == "Review"
    assert values["Calculation"] == "=LEN(A2)"


def test_invalid_url_batch_leaves_no_output(tmp_path):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    source(original)
    with pytest.raises(ValueError):
        update_workbook(original, output, ["https://example.com/ok", "file:///private"], as_of=TODAY)
    assert not output.exists()


def test_invalid_assessment_keeps_original_data_for_review(tmp_path):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    source(original, ("Leadership fit",), ("strong",))
    update_workbook(original, output, [], as_of=TODAY, offline=True)
    values = get_values(output)
    assert values["Assistant section"] == "Review"
    assert values["Leadership fit"] == "strong"
    assert "Evidence required" in values["Assistant review flags"]


@pytest.mark.parametrize("main_fails", [False, True])
def test_supplied_links_fallback_checks_identity_and_preserves_tracking(tmp_path, main_fails):
    original, output, again = (tmp_path / name for name in ("input.xlsx", "output.xlsx", "again.xlsx"))
    source(original, ("Other source links",), ("https://example.com/wrong\nhttps://example.com/right",))
    calls = []
    def fetch(url):
        calls.append(url)
        if url == "https://example.com/1":
            if main_fails:
                raise FetchError("Unavailable")
            return Page(url, "<main>Please log in</main>")
        company = "Other Co" if url.endswith("wrong") else "Original Co"
        data = {"@type": "JobPosting", "title": "Original Title", "hiringOrganization": {"name": company}, "description": "Manage engineering teams and delivery."}
        return Page(url, '<script type="application/ld+json">' + json.dumps(data) + '</script>')
    update_workbook(original, output, [], as_of=TODAY, fetcher=fetch)
    row = get_values(output)
    assert calls == ["https://example.com/1", "https://example.com/wrong", "https://example.com/right"]
    assert row["Assistant alternate URL"] == "https://example.com/right"
    assert row["Assistant description"] == "Manage engineering teams and delivery."
    assert row["Job posting"] == "https://example.com/1"
    assert row["Comment"] == "Do not edit"
    calls.clear()
    update_workbook(output, again, [], as_of=TODAY, fetcher=fetch)
    assert calls == ["https://example.com/right"]


def test_supplied_links_mismatch_stays_unverified(tmp_path):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    source(original, ("Other source links",), ("https://example.com/wrong",))
    def fetch(url):
        if url.endswith("/1"):
            return Page(url, "<main>Log in</main>")
        data = {"@type": "JobPosting", "title": "Different Role", "hiringOrganization": {"name": "Original Co"}, "description": "Different job."}
        return Page(url, '<script type="application/ld+json">' + json.dumps(data) + '</script>')
    update_workbook(original, output, [], as_of=TODAY, fetcher=fetch)
    row = get_values(output)
    assert not row["Assistant alternate URL"]
    assert "did not verify" in row["Assistant discovery notes"]


def test_location_hold_is_assessed_ranked_and_flagged(tmp_path):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    source(original, ("Can work from Halifax?", "Leadership fit", "Assessment evidence", "Working arrangement", "Office city", "Location evidence"),
           ("No - Toronto office required", "good", "Comparable SaaS engineering team leadership.", "hybrid", "Toronto", "Office attendance in Toronto required."))
    update_workbook(original, output, [], as_of=TODAY, offline=True)
    row = get_values(output)
    assert row["Assistant section"] == "Outstanding"
    assert row["Apply order"] == 1
    assert row["Assistant location eligibility"] == "assumed acceptable"
    assert "Location needs checking" in row["Assistant review flags"]
    assert "Toronto office required" in row["Assistant review flags"]
    assert row["Can work from Halifax?"] == "No - Toronto office required"


@pytest.mark.parametrize('confirmed', [True,False])
def test_later_open_confirmation_supersedes_stale_deadline(tmp_path,confirmed):
    original,output=tmp_path/'input.xlsx',tmp_path/'output.xlsx'
    source(original,('Deadline','Verification status','Last verified','Leadership fit','Assessment evidence'),
           (date(2026,9,1),'open' if confirmed else 'unverified',TODAY if confirmed else None,'good','Comparable team leadership and software delivery.'))
    update_workbook(original,output,[],as_of=TODAY,offline=True)
    row=get_values(output)
    assert row['Assistant section']==('Outstanding' if confirmed else 'Review')
    assert row['Deadline'].date()==date(2026,9,1)
    if confirmed:
        assert row['Apply order']==1
        assert 'superseded' in row['Assistant review flags']


def test_cgi_future_opportunity_can_be_assessed_without_claiming_open(tmp_path):
    original,output=tmp_path/'input.xlsx',tmp_path/'output.xlsx'
    source(original,('Vacancy status','Leadership fit','Assessment evidence'),
           ('Future opportunity - no staffed vacancy','good','Comparable project and people leadership.'))
    wb=load_workbook(original);wb.active['A2']='CGI';wb.save(original);wb.close()
    update_workbook(original,output,[],as_of=TODAY,offline=True)
    row=get_values(output)
    assert row['Apply order']==1
    assert row['Assistant section']=='Outstanding'
    assert 'Future opportunity' in row['Assistant review flags']
    assert 'not confirmed' in row['Assistant review flags']
    assert row['Vacancy status']=='Future opportunity - no staffed vacancy'

def test_protected_rows_preserve_assessments_and_relative_order(tmp_path):
    original, output, again = (tmp_path / name for name in ('input.xlsx', 'output.xlsx', 'again.xlsx'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Assistant queue'
    saved = {'Assistant assessment status': 'Complete', 'Assistant overall fit': 'good',
             'Assistant assessment summary': 'Saved assessment', 'Assistant evaluated on': '2026-01-01',
             'Assistant review flags': 'Saved flags'}
    ws.append(['Company', 'Position title', 'Job posting', 'Application status', *saved])
    entries = [('Rejected B', 'Rejected'), ('Applied B', 'Applied'), ('Interview B', 'Interview'),
               ('Closed B', 'Closed'), ('Applied A', 'Applied'), ('Interview A', 'Interviewing'),
               ('Rejected A', 'Rejected'), ('Closed A', 'Closed')]
    for name, status in entries:
        ws.append([name, 'EM', None, status, *saved.values()])
    wb.save(original)
    wb.close()
    class ForbiddenAssessment:
        def resolve(self, *args, **kwargs):
            raise AssertionError('Protected row assessed')
    def forbidden_fetch(url):
        raise AssertionError('Protected row fetched')
    for source_path, target in [(original, output), (output, again)]:
        update_workbook(source_path, target, [], as_of=TODAY, fetcher=forbidden_fetch,
                        assessment_service=ForbiddenAssessment())
        wb = load_workbook(target)
        data = list(wb['Assistant queue'].values)
        rows = [dict(zip(data[0], row)) for row in data[1:]]
        wb.close()
        assert [r['Company'] for r in rows] == ['Interview B', 'Interview A', 'Applied B', 'Applied A',
                                               'Closed B', 'Closed A', 'Rejected B', 'Rejected A']
        for row in rows:
            assert row['Apply order'] == '-'
            for key, value in saved.items():
                assert row[key] == value


def test_generic_refresh_preserves_description_and_provenance(tmp_path):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    wb = Workbook(); ws = wb.active
    ws.append(["Company", "Position title", "Job posting", "Application status", "Assistant description",
               "Assistant extracted posting date", "Assistant snapshot", "Assistant fetched on"])
    ws.title = "Assistant queue"
    description = "Lead the platform team and deliver reliable services. " * 8
    ws.append(["Example", "Manager", "https://example.com/role", "Not applied", description,
               "2026-09-01", "earlier-snapshot.json", "2026-10-01"])
    wb.save(original); wb.close()
    update_workbook(original, output, [], as_of=TODAY,
                    fetcher=lambda url: Page(url, "<main>Welcome to our careers page</main>"))
    values = get_values(output)
    assert values["Assistant description"] == description
    assert values["Assistant snapshot"] == "earlier-snapshot.json"
    assert values["Assistant fetched on"] == "2026-10-01"
    assert values["Assistant extracted posting date"] == "2026-09-01"
    assert "retained earlier role evidence" in values["Assistant review flags"]
    assert values["Application status"] == "Not applied"


@pytest.mark.parametrize("scoped", [False, True])
def test_refresh_accepts_scoped_duties_but_not_long_generic_text(tmp_path, scoped):
    original, output = tmp_path / "input.xlsx", tmp_path / "output.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "Assistant queue"
    ws.append(["Company", "Position title", "Job posting", "Application status", "Assistant description"])
    earlier = "Lead engineers and deliver platform services. " * 8
    ws.append(["Example", "Manager", "https://example.com/role", "Not applied", earlier])
    wb.save(original); wb.close()
    if scoped:
        current = "Lead the product team and manage reliable delivery. " * 3
        html = '<h1>Manager</h1><span class="jobdescription">' + current + '</span>'
    else:
        current = "Welcome to our company and learn about our community. " * 20
        html = "<main>" + current + "</main>"
    update_workbook(original, output, [], as_of=TODAY, fetcher=lambda url: Page(url, html))
    values = get_values(output)
    assert values["Assistant description"].strip() == (current if scoped else earlier).strip()
    assert values["Assistant latest fetch snapshot"]
